# -*- coding: utf-8 -*-
import pandas as pd
import logging
import base64
import io
import os
import unicodedata
import ast
from lxml import etree


def ast_literal_eval_safe(value):
    """Evalúa un dominio Odoo en formato string de forma segura."""
    if not value:
        return []
    if isinstance(value, list):
        return value
    return ast.literal_eval(value)
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime, date, timedelta

from odoo import api, fields, models, _
from odoo.tools import DEFAULT_SERVER_DATETIME_FORMAT
from odoo.exceptions import ValidationError

from odoo.addons.l10n_co_exogenous_information_reporting.tools.utils import _column_name_field
from odoo.addons.l10n_co_exogenous_information_reporting.tools.utils import _check_dv


_logger = logging.getLogger(__name__)


class L10ncoExogenousFormatSetting(models.Model):
    _name = "l10n_co.exogenous_format_setting"
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _description = "Modelo de configuración de formatos exógenos"
    _check_company_auto = True
    _rec_name = 'format_id'

    @api.model
    def _auto_init(self):
        res = super()._auto_init()
        self._cr.execute("""
            DO $$ BEGIN
                IF NOT EXISTS (SELECT 1 FROM pg_indexes WHERE indexname = 'l10n_co_exogenous_format_setting_format_id_idx') THEN
                    CREATE INDEX l10n_co_exogenous_format_setting_format_id_idx ON l10n_co_exogenous_format_setting (format_id);
                END IF;
                IF NOT EXISTS (SELECT 1 FROM pg_indexes WHERE indexname = 'l10n_co_exogenous_format_setting_company_id_idx') THEN
                    CREATE INDEX l10n_co_exogenous_format_setting_company_id_idx ON l10n_co_exogenous_format_setting (company_id);
                END IF;
            END $$;
        """)
        _logger.info("l10n_co.exogenous_format_setting: _auto_init completed with custom indexes")
        return res

    @api.constrains('is_it_with_date_range', 'date_start', 'date_end')
    def _check_dates(self):
        for record in self:
            if record.is_it_with_date_range and record.date_start and record.date_end:
                if record.date_start > record.date_end:
                    raise ValidationError(
                        _("La fecha inicial debe ser menor a la fecha final"))

                if record.date_start.year != record.date_end.year:
                    raise ValidationError(
                        _("La fecha inicial y la fecha final deben estar en el mismo año"))

    active = fields.Boolean(string='Activo', default=True, tracking=True)
    company_id = fields.Many2one(
        comodel_name='res.company', string='Compañía', default=lambda self: self.env.company.id)
    company_ids = fields.Many2many(
        comodel_name='res.company',
        relation='l10n_co_exogenous_format_setting_company_rel',
        column1='setting_id',
        column2='company_id',
        string='Sucursales / Compañías adicionales',
        help=(
            'Empresas adicionales cuyas líneas de movimiento se incluyen en el reporte. '
            'Útil cuando la compañía reportante tiene sucursales en Odoo con NIT compartido. '
            'Si está vacío, solo se usa la compañía principal.'
        ))

    def _get_reporting_company_ids(self):
        """Retorna los IDs de compañías a usar en el dominio de generación.
        Si hay sucursales configuradas, las incluye junto a la principal."""
        self.ensure_one()
        if self.company_ids:
            return (self.company_id | self.company_ids).ids
        return self.company_id.ids

    format_id = fields.Many2one('l10n_co.exogenous_format', string='Formato', tracking=True, check_company=True)
    format_setting_line_ids = fields.One2many(
        comodel_name='l10n_co.exogenous_format_setting_line', inverse_name='format_setting_id', string='Líneas de configuración')
    apply_concepts = fields.Boolean(
        string='Aplicar conceptos', related='format_id.apply_concepts', readonly=True, store=True)

    is_it_with_date_range = fields.Boolean(
        string='Rango de fechas?', related='format_id.is_it_with_date_range', readonly=True, store=True)
    date_start = fields.Date(string='Fecha inicial')
    date_end = fields.Date(string='Fecha final')

    journal_ids = fields.Many2many(
        comodel_name='account.journal',
        relation='l10n_co_exogenous_format_setting_journal_rel',
        string='Diarios a excluir')
    partner_ids = fields.Many2many(
        comodel_name='res.partner',
        relation='l10n_co_exogenous_format_setting_partner_rel',
        string='Contactos a excluir')
    tax_ids = fields.Many2many(
        comodel_name='account.tax',
        relation='l10n_co_exogenous_format_setting_tax_rel',
        string='Impuestos a excluir',
        help='Excluir líneas de movimiento que tengan estos impuestos como origen (tax_line_id)')

    use_custom_domain = fields.Boolean(
        string='Usar dominio personalizado',
        default=False,
        help='Activar modo desarrollador para escribir condiciones de exclusion avanzadas')
    colaboracion_contract_id = fields.Many2one(
        'l10n_co.exogenous_colaboracion_contract',
        string='Contrato de colaboración',
        domain="[('operator_company_id', '=', company_id)]",
        help='Solo aplica para formatos 5247-5252. Filtra los apuntes a los del contrato '
             'y rellena tcon, idfi, tdopa, nidpa.')
    custom_domain = fields.Char(
        string='Dominio personalizado',
        default='[]',
        help="Dominio Odoo para filtrar asientos. Ej: [('move_id.move_type', '=', 'entry')]\\n"
             "Variables disponibles: company_id, date_start, date_end\\n"
             "Este dominio se AGREGA a las exclusiones existentes (journals, partners, taxes)")
    custom_domain_preview = fields.Text(
        string='Vista previa del dominio',
        compute='_compute_custom_domain_preview',
        help='Muestra el dominio completo que se aplicara')

    binary_file = fields.Binary(string='Archivo Excel', tracking=False)
    binary_file_name = fields.Char(string='Nombre del archivo Excel', tracking=False)

    xml_file = fields.Binary(string='Archivo XML', tracking=False)
    xml_file_name = fields.Char(string='Nombre del archivo XML', tracking=False)

    flat_file = fields.Binary(string='Archivo plano', tracking=False)
    flat_file_name = fields.Char(string='Nombre archivo plano', tracking=False)

    report_type = fields.Selection(
        related='format_id.report_type', string='Tipo reporte', readonly=True, store=True)
    district_id = fields.Many2one(
        related='format_id.district_id', string='Distrito', readonly=True, store=True)

    uvt_value = fields.Float(
        string='Valor UVT',
        digits=(16, 2),
        default=0.0,
        help="Valor de la Unidad de Valor Tributario (UVT) para el período fiscal. "
             "Ej: 2024 = 47,065 COP, 2025 = 49,799 COP. "
             "Se usa para calcular límites en UVT en las columnas del reporte")
    send_number = fields.Integer(
        string='Número de envío',
        default=1,
        help='Número secuencial de envío para el XML')

    truncate_to_max_length = fields.Boolean(
        string='Truncar al máximo',
        default=False,
        help='Si está activo, los valores de texto que excedan la longitud máxima '
             'del campo se truncan automáticamente en vez de marcarse como error XSD')

    partner_group_field = fields.Selection([
        ('partner_id', 'Tercero directo'),
        ('commercial_partner_id', 'Tercero comercial'),
        ('parent_partner_id', 'Empresa matriz'),
    ], string='Agrupación de tercero', default='partner_id',
        help='Campo de res.partner para agrupar los terceros en el reporte. '
             'Tercero directo: usa el partner_id de la línea. '
             'Tercero comercial: agrupa por commercial_partner_id. '
             'Empresa matriz: agrupa por parent_id (fallback al propio partner).')

    concept_count = fields.Integer(
        string='Conceptos',
        compute='_compute_concept_count')
    setting_line_count = fields.Integer(
        string='Columnas configuradas',
        compute='_compute_setting_line_count')

    all_account_ids = fields.Many2many(
        comodel_name='account.account',
        string='Todas las cuentas',
        compute='_compute_all_account_ids',
        search='_search_all_account_ids',
        help='Cuentas configuradas en todas las líneas. Permite filtrar configuraciones por cuenta.')

    @api.depends('format_setting_line_ids', 'format_setting_line_ids.account_ids')
    def _compute_all_account_ids(self):
        for rec in self:
            rec.all_account_ids = rec.format_setting_line_ids.mapped('account_ids')

    @api.model
    def _search_all_account_ids(self, operator, value):
        """Permite buscar configuraciones que tengan una cuenta en alguna de sus líneas."""
        lines = self.env['l10n_co.exogenous_format_setting_line'].search(
            [('account_ids', operator, value)]
        )
        return [('format_setting_line_ids', 'in', lines.ids)]

    @api.constrains('format_id', 'company_id', 'active')
    def _check_unique_format_company(self):
        for rec in self.filtered('active'):
            duplicate = self.search([
                ('format_id', '=', rec.format_id.id),
                ('company_id', '=', rec.company_id.id),
                ('active', '=', True),
                ('id', '!=', rec.id),
            ], limit=1)
            if duplicate:
                raise ValidationError(_(
                    "Ya existe una configuración activa para el formato '%(format)s' "
                    "en la compañía '%(company)s'. Solo se permite una configuración por formato y empresa.",
                    format=rec.format_id.display_name,
                    company=rec.company_id.name,
                ))

    @api.depends('format_id')
    def _compute_concept_count(self):
        for rec in self:
            if rec.format_id and rec.format_id.apply_concepts:
                rec.concept_count = self.env['l10n_co.exogenous_concept'].search_count([
                    ('format_id', '=', rec.format_id.id),
                    ('active', '=', True)
                ])
            else:
                rec.concept_count = 0

    @api.depends('format_setting_line_ids')
    def _compute_setting_line_count(self):
        for rec in self:
            rec.setting_line_count = len(rec.format_setting_line_ids)
    
    @api.depends('custom_domain', 'use_custom_domain', 'journal_ids', 'partner_ids', 'tax_ids')
    def _compute_custom_domain_preview(self):
        """Genera vista previa del dominio completo que se aplicara"""
        for rec in self:
            if not rec.use_custom_domain or not rec.custom_domain:
                rec.custom_domain_preview = _("Usando exclusiones estandar (diarios, contactos, impuestos)")
                continue

            try:
                domain_parsed = self._parse_custom_domain(rec.custom_domain)
                exclusions = []
                if rec.journal_ids:
                    exclusions.append(f"Diarios excluidos: {len(rec.journal_ids)}")
                if rec.partner_ids:
                    exclusions.append(f"Contactos excluidos: {len(rec.partner_ids)}")
                if rec.tax_ids:
                    exclusions.append(f"Impuestos excluidos: {len(rec.tax_ids)}")

                preview = f"Dominio personalizado: {domain_parsed}"
                if exclusions:
                    preview += f"\n+ Exclusiones: {', '.join(exclusions)}"
                rec.custom_domain_preview = preview
            except Exception as e:
                rec.custom_domain_preview = _("Error en dominio: %s") % str(e)

    def _parse_custom_domain(self, domain_str):
        """Parsea y valida el dominio personalizado"""
        if not domain_str:
            return []
        try:
            domain = eval(domain_str, {"__builtins__": {}})
            if not isinstance(domain, list):
                raise ValidationError(_("El dominio debe ser una lista"))
            return domain
        except SyntaxError as e:
            raise ValidationError(_("Error de sintaxis en dominio: %s") % str(e))
        except Exception as e:
            raise ValidationError(_("Error al parsear dominio: %s") % str(e))

    def action_test_custom_domain(self):
        """Accion para probar el dominio personalizado"""
        self.ensure_one()
        if not self.use_custom_domain or not self.custom_domain:
            return {
                'type': 'ir.actions.client',
                'tag': 'display_notification',
                'params': {
                    'title': _('Sin dominio'),
                    'message': _('Active el modo desarrollador y escriba un dominio para probar'),
                    'type': 'warning',
                    'sticky': False,
                }
            }

        try:
            domain = self._parse_custom_domain(self.custom_domain)
            rc = self._get_reporting_company_ids()
            base_domain = [
                ('parent_state', '=', 'posted'),
                ('company_id', 'in' if len(rc) > 1 else '=', rc if len(rc) > 1 else rc[0]),
            ]
            test_domain = base_domain + domain
            count = self.env['account.move.line'].search_count(test_domain)

            return {
                'type': 'ir.actions.client',
                'tag': 'display_notification',
                'params': {
                    'title': _('Dominio valido'),
                    'message': _('El dominio encontraria %d lineas de asiento') % count,
                    'type': 'success',
                    'sticky': False,
                }
            }
        except Exception as e:
            return {
                'type': 'ir.actions.client',
                'tag': 'display_notification',
                'params': {
                    'title': _('Error en dominio'),
                    'message': str(e),
                    'type': 'danger',
                    'sticky': True,
                }
            }

    def _get_fields_many2one(self, fields_contact):
        """
            Función que nos permite obtener de los campo de Odoo de contacto cuales son de tipo Many2one
        """
        return fields_contact.mapped('field_odoo_id').filtered(lambda field: field.ttype == 'many2one')

    def _get_last_day_pass_year(self):
        """
            Función que nos permite obtener el ultimo día del año pasado
        """
        return date(date.today().year, 1, 1) - timedelta(days=1)

    def _dynamic_search_read(self, model, domain, fields, order=None):
        """
            Esta función nos permite hacer search_read de manera dinamica
        """
        result = self.env[model].search_read(
            domain=domain,
            fields=fields,
            order=order,
        )
        return result

    def _get_type_documents_by_format(self):
        """
            Esta función nos permite obtener de acuerdo al formato los tipos de documentos de la dian
        """
        model = 'l10n_co.exogenous_document_type'
        domain = [('document_type_table_ids', 'in',
                   self.format_id.document_type_table_id.id), ('type_document_id', '!=', False)]
        fields = ['type_document_id', 'code']
        return self._dynamic_search_read(model=model, domain=domain, fields=fields)

    def _get_or_create_consumidor_final(self):
        """Devuelve (o crea) el partner Consumidor Final con NIT 222222222 para
        la compañía actual. Se usa cuando un apunte no tiene tercero.
        """
        company = self.company_id or self.env.company
        Partner = self.env['res.partner'].sudo()
        cf = Partner.search([
            *Partner._check_company_domain(company),
            '|', ('vat', '=', '222222222'), ('vat', '=', '2222222222'),
        ], limit=1)
        if cf:
            return cf
        nit_type = self.env['l10n_latam.identification.type'].search([
            ('l10n_co_document_code', '=', 'rut')
        ], limit=1)
        vals = {
            'name': 'CUANTÍAS MENORES',
            'vat': '222222222',
            'is_company': True,
            'company_id': False,
        }
        if nit_type:
            vals['l10n_latam_identification_type_id'] = nit_type.id
        cf = Partner.create(vals)
        return cf

    def _get_columns_account_move_line_by_format(self):
        # Always include date + move_id: the partner-audit wizard needs them to
        # populate the "Asiento" and "Fecha" columns. The previous shortcut for
        # formats 1007/1008 saved a few bytes per row but emptied those columns.
        return self._get_fields_use_account_move_line() + [
            'date', 'partner_id', 'account_id', 'move_id', 'journal_id',
        ]

    def _get_accounts_by_format(self):
        """Obtiene mapeo account_id -> concept_code/field_id usando nueva estructura"""
        vals = {}
        for setting_line in self.format_setting_line_ids:
            if self.format_id.apply_concepts:
                concept = setting_line.concept_id
                field = setting_line.format_field_id
                # Usar configuración por campo si existe
                accounts = concept.get_accounts_for_report(field.id if field else None)
                for account in accounts:
                    if account.id not in vals:
                        vals[account.id] = concept.code
            else:
                field = setting_line.format_field_id
                accounts = field.get_accounts_for_report()
                for account in accounts:
                    if account.id not in vals:
                        vals[account.id] = field.id

        return vals

    def _normalice_data_dataframe_partner(self, fields_many2one, df_partners, df_type_documents, df_other_info, fields_to_clear_contact, fields_to_clear_company):
        for key, value in fields_many2one.items():
            if key == 'l10n_latam_identification_type_id':
                continue
            if key in df_partners.columns and key in df_other_info.columns:
                mapping_dict = {d['id']: d.get(
                    value, None) for element in df_other_info[key] for d in element}
                def _map_m2o(x, md=mapping_dict):
                    if isinstance(x, (list, tuple)):
                        return md.get(x[0], x[1])
                    return md.get(x, x)
                df_partners[key] = df_partners[key].apply(_map_m2o)

        mapping_dict = dict(zip(df_type_documents['type_document_id'].apply(
            lambda x: x[1]), df_type_documents['code']))

        def _map_doc_type(x):
            if isinstance(x, (list, tuple)):
                name = x[1]
                return mapping_dict.get(name, name)
            return mapping_dict.get(x, x)
        df_partners['l10n_latam_identification_type_id'] = df_partners['l10n_latam_identification_type_id'].apply(_map_doc_type)

        for index, row in df_partners.iterrows():
            fields_to_clear = fields_to_clear_company if row['is_company'] else fields_to_clear_contact
            for column in fields_to_clear:
                df_partners.at[index, column] = None

        return df_partners

    def _normalice_and_merge_data_dataframe_account_move_line(self, df_account_move_lines, df_partners):
        
        df_account_move_lines['partner_id'] = df_account_move_lines['partner_id'].apply(
            lambda x: x[0])

        # Realizar la mezcla basada en la columna partner_id y la columna id
        df_result = pd.merge(df_account_move_lines, df_partners,
                             left_on='partner_id', right_on='id', how='left')

        # Eliminar la columna id duplicada, si es necesario
        df_result.drop('id_y', axis=1, inplace=True)
        df_result.drop('id_x', axis=1, inplace=True)
        df_result.drop('partner_id', axis=1, inplace=True)
        
        return df_result

    def _get_field_name_second_field_many2one(self, fields_contact):
        vals = dict()
        for field in fields_contact.filtered(lambda field: field.ttype == 'many2one'):
            if field.field_odoo_id.name not in vals:
                vals[field.field_odoo_id.name] = field.field_odoo_internal_id.name
            else:
                vals[field.field_odoo_id.name] = list(
                    vals.get(field.field_odoo_id.name)).append(field.field_odoo_internal_id.name)
        return vals

    def _get_information_by_account_move_line(self, accounts_ids, format_field=None, setting_line=None):
        """Obtiene lineas de asiento aplicando exclusiones, filtros de impuesto y corte

        Args:
            accounts_ids: lista de IDs de cuentas contables
            format_field: campo de formato para aplicar filtros de impuesto y corte
            setting_line: línea de configuración para filtros adicionales de condición
        """
        reporting_companies = self._get_reporting_company_ids()
        company_operator = 'in' if len(reporting_companies) > 1 else '='
        company_value = reporting_companies if len(reporting_companies) > 1 else reporting_companies[0]

        domain = [
            ('parent_state', '=', 'posted'),
            ('company_id', company_operator, company_value),
            ('account_id', 'in', accounts_ids)
        ]

        # Diagnostic: how many lines exist BEFORE applying exclusions, so when
        # a user reports "exclusions don't work" they can see actual counts.
        if self.journal_ids or self.tax_ids or self.partner_ids:
            try:
                _pre_n = self.env['account.move.line'].search_count(domain)
            except Exception:
                _pre_n = -1
        else:
            _pre_n = None

        fields = self._get_columns_account_move_line_by_format()
        order = 'date asc'

        # Filtro por fechas segun corte del campo
        corte = format_field.corte if format_field else 'year_movement'
        if not self.is_it_with_date_range:
            domain.append(('date', '<=', self._get_last_day_pass_year()))
        elif corte == 'final_balance':
            domain.append(('date', '<=', self.date_end))
        else:
            domain += [('date', '>=', self.date_start), ('date', '<=', self.date_end)]

        # Exclusiones estandar
        if self.journal_ids:
            domain.append(('journal_id', 'not in', self.journal_ids.ids))

        if self.tax_ids:
            # Exclude BOTH the tax line itself (where tax_line_id is set) AND
            # the base lines that carry this tax (tax_ids many2many). The old
            # version only excluded tax_line_id, which left base lines in the
            # report — exactly what the user expects to be removed when they
            # add a tax to "Impuestos a excluir".
            domain += [
                ('tax_line_id', 'not in', self.tax_ids.ids),
                ('tax_ids', 'not in', self.tax_ids.ids),
            ]

        if self.partner_ids:
            domain += [('partner_id', 'not in', self.partner_ids.ids)]

        domain.append(('l10n_co_exogenous_skip', '=', False))
        domain.append(('move_id.l10n_co_exogenous_skip', '=', False))

        # Filtro por contrato de colaboración: si el setting tiene contrato fijado,
        # solo trae líneas de ese contrato. Si no, EXCLUYE las que tengan contrato
        # (deben reportarse en 5247-5252, no en 1001/1003/etc.).
        if self.colaboracion_contract_id:
            domain.append(('l10n_co_exogenous_colaboracion_contract_id', '=',
                           self.colaboracion_contract_id.id))
        else:
            domain.append(('l10n_co_exogenous_colaboracion_contract_id', '=', False))

        # Filtros de impuesto por campo
        if format_field:
            if format_field.line_tax_filter == 'tax_lines_only':
                domain.append(('tax_line_id', '!=', False))
            elif format_field.line_tax_filter == 'base_lines_only':
                domain.append(('tax_line_id', '=', False))

            if format_field.tax_group_ids:
                if format_field.line_tax_filter == 'tax_lines_only':
                    domain.append(('tax_line_id.tax_group_id', 'in', format_field.tax_group_ids.ids))
                else:
                    domain.append(('tax_ids.tax_group_id', 'in', format_field.tax_group_ids.ids))

            if format_field.tax_ids_filter:
                if format_field.line_tax_filter == 'tax_lines_only':
                    domain.append(('tax_line_id', 'in', format_field.tax_ids_filter.ids))
                else:
                    domain.append(('tax_ids', 'in', format_field.tax_ids_filter.ids))

            # Dominio adicional configurable desde la UI (widget="domain")
            if getattr(format_field, 'filter_domain', None):
                try:
                    extra = ast_literal_eval_safe(format_field.filter_domain)
                    if isinstance(extra, list):
                        domain += extra
                except Exception:
                    _logger.warning(
                        'filter_domain inválido en format_field id=%s: %r',
                        format_field.id, format_field.filter_domain)

        # Filtros de condición por línea (F2)
        if setting_line:
            if setting_line.move_type_filter and setting_line.move_type_filter != 'all':
                domain.append(('move_id.move_type', '=', setting_line.move_type_filter))
            if setting_line.exclude_reconciled:
                domain.append(('full_reconcile_id', '=', False))
            if setting_line.only_with_tax:
                domain += ['|', ('tax_ids', '!=', False), ('tax_line_id', '!=', False)]
            if setting_line.journal_ids:
                domain.append(('journal_id', 'in', setting_line.journal_ids.ids))
            if setting_line.custom_line_domain:
                try:
                    custom_line = self._parse_custom_domain(setting_line.custom_line_domain)
                    domain += custom_line
                except (ValidationError, Exception) as e:
                    _logger.warning("Error en dominio de línea %s: %s", setting_line.id, str(e))

        # Aplicar dominio personalizado (modo desarrollador)
        if self.use_custom_domain and self.custom_domain:
            try:
                custom = self._parse_custom_domain(self.custom_domain)
                domain += custom
                _logger.info("Dominio personalizado aplicado: %s", custom)
            except ValidationError as e:
                _logger.warning("Error en dominio personalizado: %s", str(e))

        partner_ids = set()
        account_move_lines = self._dynamic_search_read(
            model='account.move.line', domain=domain, fields=fields, order=order)

        # Para líneas sin tercero, usar el "Consumidor Final" (NIT 222222222)
        # — DIAN exige tercero. Si no existe se crea on-the-fly (1 vez por compañía).
        cf = self._get_or_create_consumidor_final()
        cf_tuple = [cf.id, cf.display_name] if cf else None
        if cf_tuple:
            for aml in account_move_lines:
                if not aml.get('partner_id'):
                    aml['partner_id'] = list(cf_tuple)

        if _pre_n is not None:
            _logger.info(
                "[exo exclusiones] line=%s field=%s — antes=%d/desp=%d (excluidas=%d) "
                "[journals=%d, taxes=%d, partners=%d]",
                setting_line.id if setting_line else '-',
                format_field.name if format_field else '-',
                _pre_n, len(account_move_lines), _pre_n - len(account_move_lines),
                len(self.journal_ids), len(self.tax_ids), len(self.partner_ids))

        # Centralizado: si el diario tiene `l10n_co_exogenous_partner_id`, ese
        # tercero se usa SIEMPRE (sin importar group_by_journal). Caso típico:
        # diarios bancarios con el banco como tercero, diarios de nómina con
        # el ente recaudador, etc. Si el diario NO tiene tercero configurado,
        # se preserva el partner del apunte como antes.
        if account_move_lines:
            journal_ids_in_lines = {aml['journal_id'][0] for aml in account_move_lines if aml.get('journal_id')}
            if journal_ids_in_lines:
                journals = self.env['account.journal'].browse(list(journal_ids_in_lines))
                jpartner_by_journal = {
                    j.id: (j.l10n_co_exogenous_partner_id.id, j.l10n_co_exogenous_partner_id.display_name)
                    for j in journals if j.l10n_co_exogenous_partner_id
                }
                if jpartner_by_journal:
                    for aml in account_move_lines:
                        jid = aml['journal_id'][0] if aml.get('journal_id') else None
                        if jid and jid in jpartner_by_journal:
                            aml['partner_id'] = list(jpartner_by_journal[jid])

        for aml in account_move_lines:
            if aml['partner_id']:
                partner_ids.add(aml['partner_id'][0])

        return partner_ids, account_move_lines

    def _resolve_partner_grouping(self, move_lines, partner_ids):
        """Remapea partner_id según partner_group_field configurado.

        Args:
            move_lines: lista de dicts con campo partner_id [id, name]
            partner_ids: set de IDs de partners originales

        Returns:
            (move_lines_remapped, new_partner_ids)
        """
        group_field = self.partner_group_field or 'partner_id'
        if group_field == 'partner_id' or not partner_ids:
            return move_lines, partner_ids

        # Leer el campo alterno de todos los partners involucrados
        partners = self.env['res.partner'].browse(list(partner_ids))
        mapping = {}  # original_partner_id -> (new_partner_id, new_partner_name)

        if group_field == 'commercial_partner_id':
            for p in partners:
                cp = p.commercial_partner_id
                if cp:
                    mapping[p.id] = (cp.id, cp.display_name)
                else:
                    mapping[p.id] = (p.id, p.display_name)
        elif group_field == 'parent_partner_id':
            for p in partners:
                parent = p.parent_id
                if parent:
                    mapping[p.id] = (parent.id, parent.display_name)
                else:
                    mapping[p.id] = (p.id, p.display_name)

        if not mapping:
            return move_lines, partner_ids

        new_partner_ids = set()
        for aml in move_lines:
            pid = aml.get('partner_id')
            if pid and isinstance(pid, (list, tuple)):
                original_id = pid[0]
                new_id, new_name = mapping.get(original_id, (original_id, pid[1]))
                aml['partner_id'] = [new_id, new_name]
                new_partner_ids.add(new_id)
            elif pid:
                new_partner_ids.add(pid)

        return move_lines, new_partner_ids

    def _get_values_form_many2one(self, data_source, fields_many2one):
        """ 
            Funcion que nos permite de acuerdo a una informacion obtenida por un search_read
            y con base a unos campos que son de tipo many2one obtener los ids de esos modelos
            y llevarlos en una lista
        """
        values_fields_many2one = dict()
        for ds in data_source:
            for field_many2one in fields_many2one:
                if not isinstance(ds[field_many2one], bool):
                    if field_many2one not in values_fields_many2one:
                        values_fields_many2one[field_many2one] = [
                            ds[field_many2one][0]]
                    else:
                        values_fields_many2one[field_many2one].append(
                            ds[field_many2one][0])
        return values_fields_many2one

    def _get_values_form_dynamic_search_read(self, search_read_by_dynamic_model, field_odoo_internal):
        values = list()
        for value in search_read_by_dynamic_model:
            values.append(value[field_odoo_internal.name])
        return values

    def _get_data_from_dynamic_model(self, partner_info, fields_contact):

        fields_many2one_odoo = self._get_fields_many2one(fields_contact)
        fields_names_many2one_odoo = fields_many2one_odoo.mapped("name")
        values_fields_many2one = self._get_values_form_many2one(
            partner_info, fields_names_many2one_odoo)

        list_fields_name_and_models = fields_contact.mapped('field_odoo_id').filtered(
            lambda field: field.ttype == 'many2one').mapped(lambda field: {field.name: field.relation})
        values_by_dynamic_models = dict()

        for dict_field_name_and_model in list_fields_name_and_models:
            for values in dict_field_name_and_model:
                if values_fields_many2one.get(values, False):
                    field_odoo_internal = fields_contact.filtered(
                        lambda field_line: field_line.field_odoo_id.name == values).mapped('field_odoo_internal_id')

                    model = dict_field_name_and_model[values]
                    domain = [
                        ('id', 'in', list(set(values_fields_many2one.get(values))))]
                    fields = field_odoo_internal.mapped('name')
                    search_read_by_dynamic_model = self._dynamic_search_read(
                        model=model, domain=domain, fields=fields)

                    if values not in values_by_dynamic_models:
                        values_by_dynamic_models[values] = search_read_by_dynamic_model
                    else:
                        values_by_dynamic_models[values].extend(
                            search_read_by_dynamic_model)

        return values_by_dynamic_models

    def _get_information_partner(self, partner_ids, fields_contact):

        if not fields_contact:
            return self._show_message_error(f"No hay campos configurados para la compañía: {self.company_id.name}")

        fields_odoo = fields_contact.mapped('field_odoo_id.name')
        fields_odoo.append('is_company')

        model = 'res.partner'
        domain = [('id', 'in', list(partner_ids))]
        fields = fields_odoo

        partner_info = self._dynamic_search_read(
            model=model, domain=domain, fields=fields)

        other_info_contact = self._get_data_from_dynamic_model(
            partner_info, fields_contact)
        return partner_info, other_info_contact

    def _get_accounts(self, format_setting_line):
        """Obtiene cuentas: primero de la línea, luego del concepto si aplica"""
        if format_setting_line.account_ids:
            return format_setting_line.account_ids
        if self.format_id.apply_concepts and format_setting_line.concept_id:
            concept = format_setting_line.concept_id
            field = format_setting_line.format_field_id
            return concept.get_accounts_for_report(field.id if field else None)
        return self.env['account.account']

    def _get_field_concept(self):
        return self.env.ref('l10n_co_exogenous_information_reporting.l10n_co_exogenous_format_field_cpt', raise_if_not_found=True)

    def _get_columns_ordered(self, format_fields) -> list:
        vals = list()
        for format_field in format_fields:
            if format_field.id == self._get_field_concept().id:
                vals.append('account_id')
            elif format_field.field_odoo_id:
                vals.append(format_field.field_odoo_id.name)
        return vals

    def _get_fields_to_clear_contact(self, fields_contact):
        return list(set(fields_contact.mapped('field_odoo_id').mapped('name')) - set(fields_contact.filtered(lambda field: field.source == 'contact' and field.applies_to_contact).mapped('field_odoo_id').mapped('name')))

    def _get_fields_to_clear_company(self, fields_contact):
        return list(set(fields_contact.mapped('field_odoo_id').mapped('name')) - set(fields_contact.filtered(lambda field: field.source == 'contact' and field.applies_to_company).mapped('field_odoo_id').mapped('name')))

    def find_column_position(self, sheet, column_name):
        for col_idx in range(1, sheet.max_column + 1):
            if sheet.cell(row=1, column=col_idx).value == column_name:
                return col_idx
        return None  # Retorna None si no se encuentra la columna

    def _format_field_by_accumulated(self, format_setting_line):
        """Obtiene la fórmula de acumulación usando configuración por campo

        Mapea cada cuenta a la fórmula configurada (DB, CR, DB-CR, CR-DB, bases, etc.)
        para que get_column_value() pueda calcular el valor correcto.
        """
        vals = dict()
        field = format_setting_line.format_field_id
        accounts = self._get_accounts(format_setting_line)

        # Naturaleza: concepto (específico por campo) > línea > default
        # El concepto tiene configuración por campo (ej. pago=db_cr, retp=cr_db)
        # que es más específica que el default genérico de la línea.
        nature = None
        if self.format_id.apply_concepts and format_setting_line.concept_id:
            concept_nature = format_setting_line.concept_id.get_nature_for_report(field.id if field else None)
            if concept_nature and concept_nature != 'db_cr':
                nature = concept_nature
        if not nature:
            nature = format_setting_line.nature_account
        nature = nature or 'db_cr'

        accounts_accumulated_by = dict()
        for account in accounts:
            accounts_accumulated_by[(account.id, account.display_name)] = nature
        vals[field.name] = accounts_accumulated_by

        return vals

    def _get_fields_use_account_move_line(self):
        return ['credit', 'debit', 'balance', 'tax_base_amount']

    @staticmethod
    def _compute_formula(row, formula):
        """Calcula el valor según la fórmula DIAN configurada

        Args:
            row: dict o Series con campos debit, credit, balance, tax_base_amount
            formula: string con la fórmula (debit, credit, db_cr, cr_db, base_calc_*, tax_base_amount)
        """
        debit = row.get('debit', 0) or 0
        credit = row.get('credit', 0) or 0
        tax_base = row.get('tax_base_amount', 0) or 0

        if formula == 'debit':
            return debit
        elif formula == 'credit':
            return credit
        elif formula == 'db_cr':
            return debit - credit
        elif formula == 'cr_db':
            return credit - debit
        elif formula == 'base_calc_db':
            return tax_base if debit > 0 else 0
        elif formula == 'base_calc_cr':
            return tax_base if credit > 0 else 0
        elif formula == 'base_calc_db_cr':
            return tax_base if (debit - credit) != 0 else 0
        elif formula == 'base_calc_cr_db':
            return tax_base if (credit - debit) != 0 else 0
        elif formula == 'tax_base_amount':
            return tax_base
        # Compatibilidad con valores antiguos
        elif formula == 'balance':
            return row.get('balance', 0) or 0

        return row.get('balance', 0) or 0

    def _get_fields_format_id(self, format_setting_line):
        return format_setting_line.mapped('format_field_id').mapped('name')

    def _create_original_dataframe(self, format_fields):
        return pd.DataFrame(columns=format_fields.mapped('name'))

    def _get_fields_odoo_and_format(self, fields_contact):
        vals = dict()
        if self.format_id.apply_concepts:
            vals['account_id'] = self._get_field_concept().name

        for field_contact in fields_contact: 
            vals[field_contact.field_odoo_id.name] = field_contact.name

        return vals

    def _get_fields_fill_smaller_amount(self):
        return {
            self.env.ref('l10n_co_exogenous_information_reporting.l10n_co_exogenous_format_field_tdoc', raise_if_not_found=True).name : '43',
            self.env.ref('l10n_co_exogenous_information_reporting.l10n_co_exogenous_format_field_nid', raise_if_not_found=True).name : '222222222',
            self.env.ref('l10n_co_exogenous_information_reporting.l10n_co_exogenous_format_field_raz', raise_if_not_found=True).name : 'cuantías menores',
        }
    

    def _generate_row_by_smaller_amount(self, row_smaller_amount, df):
        new_row = pd.DataFrame(row_smaller_amount, index=[0])
        new_row_columns = df.columns
        new_row = new_row.reindex(columns=new_row_columns)
        df = pd.concat([df, new_row], ignore_index=True)
        return df



    @staticmethod
    def _validate_xsd_value(value, format_field):
        """Valida un valor contra las restricciones XSD del campo.

        Returns:
            tuple: (is_valid: bool, errors: list[str])
        """
        import re
        errors = []
        if value is None or (isinstance(value, str) and not value.strip()):
            if format_field.xsd_required:
                errors.append("Requerido")
            return (len(errors) == 0, errors)

        str_val = str(value).strip()

        # Detectar valores negativos en campos numéricos
        if format_field.source == 'journal_items':
            try:
                num_check = float(str_val)
                if num_check < 0:
                    errors.append(f"Valor negativo: {str_val}")
            except (ValueError, TypeError):
                pass

        # Validar longitud
        if format_field.min_length and len(str_val) < format_field.min_length:
            errors.append(f"Min largo {format_field.min_length}, tiene {len(str_val)}")
        if format_field.max_length and len(str_val) > format_field.max_length:
            errors.append(f"Max largo {format_field.max_length}, tiene {len(str_val)}")

        # Validar patrón
        if format_field.xsd_pattern:
            try:
                if not re.fullmatch(format_field.xsd_pattern, str_val):
                    errors.append(f"No cumple patrón {format_field.xsd_pattern}")
            except re.error:
                pass

        # Validar rango numérico
        if format_field.xsd_type in ('int', 'long', 'double'):
            try:
                num_val = float(str_val)
                if format_field.xsd_min_value:
                    if num_val < float(format_field.xsd_min_value):
                        errors.append(f"Min valor {format_field.xsd_min_value}")
                if format_field.xsd_max_value:
                    if num_val > float(format_field.xsd_max_value):
                        errors.append(f"Max valor {format_field.xsd_max_value}")
            except (ValueError, TypeError):
                errors.append(f"Se esperaba tipo {format_field.xsd_type}")

        return (len(errors) == 0, errors)

    @staticmethod
    def _get_xsd_restriction_summary(format_field):
        """Genera un resumen legible de las restricciones XSD del campo."""
        parts = []
        if format_field.xsd_type and format_field.xsd_type != 'string':
            parts.append(f"tipo:{format_field.xsd_type}")
        if format_field.xsd_pattern:
            parts.append(f"patrón:{format_field.xsd_pattern}")
        if format_field.min_length:
            parts.append(f"minLen:{format_field.min_length}")
        if format_field.max_length:
            parts.append(f"maxLen:{format_field.max_length}")
        if format_field.xsd_min_value:
            parts.append(f"min:{format_field.xsd_min_value}")
        if format_field.xsd_max_value:
            parts.append(f"max:{format_field.xsd_max_value}")
        if format_field.xsd_required:
            parts.append("REQUERIDO")
        return " | ".join(parts) if parts else ""

    def generate_and_download_report(self):
        # Garantizar contexto ORM correcto: empresa principal + sucursales si aplica
        all_companies = self.company_id | self.company_ids
        self = self.with_company(all_companies)
        # Check if format settings are available
        if not self.format_setting_line_ids:
            return self._show_message_error(f"No hay configuración de formato para el formato: {self.format_id.code}")

        # Retrieve format fields
        format_fields = self.env['l10n_co.exogenous_format_field'].search(
            [('format_ids', 'in', self.format_id.id)], order="sequence asc")

        if not format_fields:
            return self._show_message_error(f"No hay campos configurados para el formato: {self.format_id.code}")

        # Filter fields based on source
        fields_contact = format_fields.filtered(lambda ff: ff.source == 'contact')

        # Get fields to clear for contact and company
        fields_to_clear_contact = self._get_fields_to_clear_contact(fields_contact)
        fields_to_clear_company = self._get_fields_to_clear_company(fields_contact)

        # Get unique keys for format
        unique_keys_by_format = format_fields.filtered(lambda ff: ff.is_unique_key).mapped(
            'field_odoo_id').mapped('name')

        unique_keys_by_format_field = format_fields.filtered(lambda ff: ff.is_unique_key).mapped('name')

        if self.format_id.apply_concepts:
            unique_keys_by_format = unique_keys_by_format + ['account_id']

        wb = Workbook()
        ws = wb.active
        ws = _column_name_field(format_fields.mapped('name'), ws)

        # Create original dataframe
        original_dataframe = self._create_original_dataframe(format_fields)
        columnas_originales = original_dataframe.columns
        row_smaller_amount = self._get_fields_fill_smaller_amount()
        concepts = self._get_accounts_by_format()

        get_information = list()

        for setting_line in self.format_setting_line_ids:
            accounts_ids = self._get_accounts(setting_line).ids
            partner_ids, account_move_lines = self._get_information_by_account_move_line(accounts_ids, format_field=setting_line.format_field_id, setting_line=setting_line)
            account_move_lines = self._apply_manual_column_overrides(account_move_lines, setting_line.format_field_id)
            account_move_lines, partner_ids = self._resolve_partner_grouping(account_move_lines, partner_ids)

            _logger.info(setting_line)
            if not account_move_lines or not partner_ids:
                get_information.append(False)
                continue
            else:
                get_information.append(True)

            partners, other_info_contact = self._get_information_partner(partner_ids, fields_contact)

            if isinstance(partners, dict) and partners.get('tag', False):
                return partners

            df_type_documents = pd.DataFrame(self._get_type_documents_by_format())
            df_account_move_lines = pd.DataFrame(account_move_lines)
            df_partners = pd.DataFrame(partners)
            df_partners = df_partners.fillna('')
            df_partners = df_partners.replace(False, '')
            df_other_info = pd.DataFrame([other_info_contact])

            df_partners = self._normalice_data_dataframe_partner(
                self._get_field_name_second_field_many2one(fields_contact),
                df_partners, df_type_documents, df_other_info, fields_to_clear_contact, fields_to_clear_company
            )

            df_account_move_lines = self._normalice_and_merge_data_dataframe_account_move_line(
                df_account_move_lines, df_partners)

            format_field_accumulated = self._format_field_by_accumulated(setting_line)

            def get_column_value(row, key):
                """Calcula el valor de la columna según la fórmula configurada"""
                column_to_get = format_field_accumulated.get(key, None)
                if isinstance(column_to_get, dict):
                    account_id = row['account_id']
                    formula = column_to_get.get(account_id, None)
                    if formula:
                        return self._compute_formula(row, formula)
                return 0.0

            columns_ordered = self._get_columns_ordered(format_fields)
            columns_ordered = columns_ordered + [setting_line.format_field_id.name]

            clave_diccionario, valor_diccionario = next(iter(format_field_accumulated.items()))

            df_account_move_lines[clave_diccionario] = df_account_move_lines.apply(
                lambda row: get_column_value(row, clave_diccionario), axis=1)

            default_concept_code = (
                setting_line.concept_id.code
                if (self.format_id.apply_concepts and setting_line.concept_id)
                else ''
            )
            df_account_move_lines['account_id'] = df_account_move_lines['account_id'].apply(
                lambda x: concepts.get(x[0], default_concept_code))

            operations = {column: 'first' if column not in self._get_fields_format_id(setting_line) else 'sum'
                        for column in columns_ordered}

            df_account_move_lines_copy = df_account_move_lines.copy(deep=True)
            filtered_grouped = df_account_move_lines_copy.groupby(unique_keys_by_format)

            result = filtered_grouped.agg(operations).round(0)
            result = result.rename(columns=self._get_fields_odoo_and_format(fields_contact))

            if self.format_id.applying_smaller_amounts:
                # Compare on absolute value: a credit-side concept (Ingresos)
                # configured with `db_cr` produces negative numbers; without
                # abs(), every value < threshold gets dropped.
                col = setting_line.format_field_id.name
                _abs = result[col].abs()
                smaller_amounts = result.loc[_abs < self.format_id.smaller_ammounts, col].sum()
                result = result.loc[_abs >= self.format_id.smaller_ammounts]


                if self.format_id.apply_concepts and smaller_amounts > 0:
                    row_smaller_amount = self._get_fields_fill_smaller_amount()
                    row_smaller_amount.update({setting_line.format_field_id.name: smaller_amounts})
                    row_smaller_amount.update({self._get_field_concept().name: setting_line.concept_id.code})
                    result = self._generate_row_by_smaller_amount(row_smaller_amount, result)
                elif not self.format_id.apply_concepts and smaller_amounts > 0:
                    row_smaller_amount.update({setting_line.format_field_id.name: smaller_amounts})

            result = result.reset_index(drop=True)
            duplicados = original_dataframe.merge(result, on=unique_keys_by_format_field, how='inner')
            filas_duplicadas = result[result.index.isin(duplicados.index)]
            for index, fila_duplicada in filas_duplicadas.iterrows():
                filtro = (original_dataframe[unique_keys_by_format_field] == fila_duplicada[unique_keys_by_format_field]).all(axis=1)
                original_dataframe.loc[filtro, fila_duplicada.index] = fila_duplicada.values

            filas_nuevas = result[~result.index.isin(duplicados.index)]
            original_dataframe = pd.concat([original_dataframe, filas_nuevas], ignore_index=True)


        if not any(get_information):
            return self._show_message_error("No se encontró información con estos parámetros")

        if not self.format_id.apply_concepts and self.format_id.applying_smaller_amounts:
            original_dataframe = self._generate_row_by_smaller_amount(row_smaller_amount, original_dataframe)

        original_dataframe = original_dataframe.reindex(columns=columnas_originales)

        # Merge final por (unique_keys + concepto): combina filas que la lógica
        # incremental de arriba dejó duplicadas. Una fila por partner/concepto.
        unique_field_names = format_fields.filtered(
            lambda f: f.is_unique_key).mapped('name')
        if self.format_id.apply_concepts:
            cpt_field = self._get_field_concept()
            if cpt_field and cpt_field.name not in unique_field_names:
                unique_field_names = unique_field_names + [cpt_field.name]
        journal_field_names = format_fields.filtered(
            lambda f: f.source == 'journal_items').mapped('name')
        group_keys = [c for c in unique_field_names if c in original_dataframe.columns]
        if group_keys:
            agg = {}
            for c in original_dataframe.columns:
                if c in group_keys:
                    continue
                if c in journal_field_names:
                    agg[c] = 'sum'
                else:
                    agg[c] = 'first'
            if agg:
                fillna_cols = {c: 0 for c in journal_field_names if c in original_dataframe.columns}
                if fillna_cols:
                    original_dataframe = original_dataframe.fillna(value=fillna_cols)
                original_dataframe = original_dataframe.groupby(
                    group_keys, as_index=False, dropna=False).agg(agg)
                original_dataframe = original_dataframe.reindex(columns=columnas_originales)

        # DIAN espera ceros, no celdas vacías. Los huecos aparecen cuando un
        # tercero tiene movimientos en una columna pero no en otra (ej: ingresos
        # sí, devoluciones no) → la concatenación de DataFrames los marca NaN.
        # Solo rellenamos las columnas de monto (source='journal_items'); las
        # de contacto (nombre, ciudad, etc.) se dejan como están.
        amount_fields = format_fields.filtered(lambda f: f.source == 'journal_items').mapped('name')
        for col in amount_fields:
            if col in original_dataframe.columns:
                original_dataframe[col] = original_dataframe[col].fillna(0)

        # Columnas derivadas y transformaciones de porcentaje/limite
        uvt = self.uvt_value or 0
        journal_fields = format_fields.filtered(lambda f: f.source == 'journal_items')
        for field in journal_fields:
            if field.name not in original_dataframe.columns:
                continue
            if field.formula_source_field_id:
                source_col = field.formula_source_field_id.name
                if source_col in original_dataframe.columns:
                    original_dataframe[field.name] = original_dataframe[source_col].apply(
                        lambda v: field.apply_value_transformations(v, uvt_value=uvt)
                        if isinstance(v, (int, float)) else v)
            elif field.formula_percentage != 100.0 or field.formula_limit_type != 'none':
                original_dataframe[field.name] = original_dataframe[field.name].apply(
                    lambda v: field.apply_value_transformations(v, uvt_value=uvt)
                    if isinstance(v, (int, float)) else v)

        def procesar_identificacion(row):
            nit = row[field_identification_number_name]

            # Validar que nit no esté vacío o sea None
            if not nit or (isinstance(nit, str) and len(nit.strip()) == 0):
                return pd.Series([nit, None])

            # Convertir a string y limpiar espacios
            nit = str(nit).strip()

            if '-' in nit:
                numero, dv_actual = nit.split('-', 1)
            elif len(nit) > 1:
                numero, dv_actual = nit[:-1], nit[-1]
            else:
                # Si solo tiene 1 carácter, no hay DV
                return pd.Series([nit, None])

            dv_nuevo = _check_dv(numero) if 'Digito de Verificación' in original_dataframe.columns else None
            return pd.Series([numero, dv_nuevo if dv_nuevo != dv_actual else dv_actual])

        filter = original_dataframe[self.env.ref('l10n_co_exogenous_information_reporting.l10n_co_exogenous_format_field_tdoc', raise_if_not_found=True).name] == '31'
        field_identification_number_name = self.env.ref('l10n_co_exogenous_information_reporting.l10n_co_exogenous_format_field_nid', raise_if_not_found=True).name
        field_dv_name = self.env.ref('l10n_co_exogenous_information_reporting.l10n_co_exogenous_format_field_dv', raise_if_not_found=True).name if 'Digito de Verificación' in original_dataframe.columns else None

        filtered_dataframe = original_dataframe[filter].reset_index()
        results = filtered_dataframe.apply(procesar_identificacion, axis=1)
        if field_dv_name:
            original_dataframe.loc[filter, [field_identification_number_name, field_dv_name]] = results.values
        else:
            original_dataframe.loc[filter, field_identification_number_name] = results.iloc[:, 0].values

        # ── Sanitizar columnas numéricas antes de validación XSD ──
        # nan → 0, float con .0 → int, negativos se mantienen para marcar como error XSD
        import math
        numeric_fields = format_fields.filtered(
            lambda f: f.source == 'journal_items'
                      and f.xsd_pattern
                      and '[0-9]' in (f.xsd_pattern or ''))
        for nf in numeric_fields:
            col = nf.name
            if col not in original_dataframe.columns:
                continue
            def _sanitize_numeric(v):
                if v is None or v is False:
                    return 0
                if isinstance(v, float):
                    if math.isnan(v) or math.isinf(v):
                        return 0
                    return int(round(v))
                if isinstance(v, str):
                    v = v.strip()
                    if not v or v.lower() == 'nan':
                        return 0
                    try:
                        fv = float(v)
                        if math.isnan(fv) or math.isinf(fv):
                            return 0
                        return int(round(fv))
                    except (ValueError, TypeError):
                        return v
                return v
            original_dataframe[col] = original_dataframe[col].apply(_sanitize_numeric)

        # ── Limpiar depto/ciudad para no-Colombia (dejar vacío) ──
        pais_ref = self.env.ref(
            'l10n_co_exogenous_information_reporting.l10n_co_exogenous_format_field_pais',
            raise_if_not_found=False)
        dpto_ref = self.env.ref(
            'l10n_co_exogenous_information_reporting.l10n_co_exogenous_format_field_dpto',
            raise_if_not_found=False)
        mun_ref = self.env.ref(
            'l10n_co_exogenous_information_reporting.l10n_co_exogenous_format_field_mun',
            raise_if_not_found=False)
        if pais_ref and pais_ref.name in original_dataframe.columns:
            pais_col = pais_ref.name
            no_co_mask = ~original_dataframe[pais_col].astype(str).str.strip().isin(['169', 'CO'])
            if dpto_ref and dpto_ref.name in original_dataframe.columns:
                original_dataframe.loc[no_co_mask, dpto_ref.name] = ''
            if mun_ref and mun_ref.name in original_dataframe.columns:
                original_dataframe.loc[no_co_mask, mun_ref.name] = ''

        # Aplicar recorte (xsd_slice) ANTES del truncate
        for ff in format_fields:
            if ff.xsd_slice and ff.xsd_slice != 'none' and ff.name in original_dataframe.columns:
                original_dataframe[ff.name] = original_dataframe[ff.name].apply(
                    lambda v, f=ff: f.apply_slice(v))

        # ── Truncar valores al max_length si la opción está activa ──
        if self.truncate_to_max_length:
            for ff in format_fields:
                if ff.max_length and ff.name in original_dataframe.columns:
                    original_dataframe[ff.name] = original_dataframe[ff.name].apply(
                        lambda v: str(v)[:ff.max_length] if v is not None and not isinstance(v, (int, float)) and len(str(v)) > ff.max_length else v)

        # Fila 2: restricciones XSD como referencia
        xsd_row = []
        for field in format_fields:
            xsd_row.append(self._get_xsd_restriction_summary(field))
        ws.append(xsd_row)

        # Estilos
        red_fill = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')
        grey_fill = PatternFill(start_color='F0F0F0', end_color='F0F0F0', fill_type='solid')
        small_font = Font(size=8, italic=True, color='666666')

        # Aplicar estilo a fila de restricciones (fila 2)
        for col_idx in range(1, len(xsd_row) + 1):
            cell = ws.cell(row=2, column=col_idx)
            cell.fill = grey_fill
            cell.font = small_font
            cell.alignment = Alignment(wrap_text=True)

        # Escribir datos con validación XSD
        field_list = list(format_fields)
        validation_errors = []  # [(fila, columna, valor, errores)]

        # Índices de columnas de depto/ciudad/país para skip en no-Colombia
        _pais_col_idx = None
        _dpto_col_idx = None
        _mun_col_idx = None
        for _fi, _ff in enumerate(field_list):
            if pais_ref and _ff.id == pais_ref.id:
                _pais_col_idx = _fi
            if dpto_ref and _ff.id == dpto_ref.id:
                _dpto_col_idx = _fi
            if mun_ref and _ff.id == mun_ref.id:
                _mun_col_idx = _fi

        for row_idx, row_data in enumerate(dataframe_to_rows(original_dataframe, index=False, header=False), start=3):
            ws.append(row_data)
            excel_row = ws.max_row

            # Determinar si esta fila es Colombia
            _is_co = True
            if _pais_col_idx is not None and _pais_col_idx < len(row_data):
                pais_val = str(row_data[_pais_col_idx]).strip()
                _is_co = pais_val in ('CO', '169')

            for col_idx, value in enumerate(row_data):
                if col_idx < len(field_list):
                    field = field_list[col_idx]
                    # Skip validación depto/ciudad si no es Colombia
                    if not _is_co and col_idx in (_dpto_col_idx, _mun_col_idx):
                        continue
                    if field.xsd_type or field.xsd_pattern or field.xsd_required:
                        is_valid, errs = self._validate_xsd_value(value, field)
                        if not is_valid:
                            ws.cell(row=excel_row, column=col_idx + 1).fill = red_fill
                            validation_errors.append((
                                excel_row, field.name, str(value) if value else '', '; '.join(errs)
                            ))

        # Hoja de errores
        if validation_errors:
            ws_err = wb.create_sheet('Errores XSD')
            ws_err.append(['Fila', 'Columna', 'Valor', 'Error'])
            header_fill = PatternFill(start_color='FF6666', end_color='FF6666', fill_type='solid')
            header_font = Font(bold=True, color='FFFFFF')
            for col_idx in range(1, 5):
                cell = ws_err.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
            for err in validation_errors:
                ws_err.append(list(err))
            # Ajustar anchos
            ws_err.column_dimensions['A'].width = 8
            ws_err.column_dimensions['B'].width = 30
            ws_err.column_dimensions['C'].width = 25
            ws_err.column_dimensions['D'].width = 50

        output = io.BytesIO()
        wb.save(output)
        file_name = f"exogena_{self.format_id.code}_{datetime.now().strftime(DEFAULT_SERVER_DATETIME_FORMAT)}"
        self.binary_file = base64.b64encode(output.getvalue())
        self.binary_file_name = f"{file_name}.xlsx"
        # Trigger an immediate browser download. Without this, the form just
        # silently refreshes and the user has to scroll to the (now-visible)
        # binary widget to download manually — looks like nothing happened.
        # Same URL pattern as generate_and_download_xml uses, which is the
        # canonical /web/content endpoint for binary fields in v14.
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/?model=%s&id=%d&field=binary_file&filename_field=binary_file_name&download=true' % (
                self._name, self.id),
            'target': 'self',
        }


    def _show_message_error(self, message):
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'type': 'danger',
                'sticky': True,
                'message': _(message),
            }
        }

    def action_load_default_columns(self):
        """Carga las columnas por defecto basadas en el formato seleccionado"""
        self.ensure_one()
        if not self.format_id:
            return self._show_message_error("Debe seleccionar un formato primero")

        FormatField = self.env['l10n_co.exogenous_format_field']
        Concept = self.env['l10n_co.exogenous_concept']
        SettingLine = self.env['l10n_co.exogenous_format_setting_line']

        format_fields = FormatField.search([
            ('format_ids', 'in', self.format_id.id),
            ('source', '=', 'journal_items')
        ], order='sequence asc')

        if not format_fields:
            return self._show_message_error(
                f"No hay campos de tipo 'journal_items' configurados para el formato {self.format_id.code}"
            )

        existing_combinations = set()
        for line in self.format_setting_line_ids:
            if self.format_id.apply_concepts:
                key = (line.concept_id.id, line.format_field_id.id)
            else:
                key = (False, line.format_field_id.id)
            existing_combinations.add(key)

        lines_created = 0

        if self.format_id.apply_concepts:
            concepts = Concept.search([
                ('format_id', '=', self.format_id.id),
                ('active', '=', True)
            ], order='code asc')

            if not concepts:
                return self._show_message_error(
                    f"No hay conceptos configurados para el formato {self.format_id.code}"
                )

            for concept in concepts:
                for field in format_fields:
                    key = (concept.id, field.id)
                    if key not in existing_combinations:
                        SettingLine.create({
                            'format_setting_id': self.id,
                            'concept_id': concept.id,
                            'format_field_id': field.id,
                        })
                        lines_created += 1
        else:
            for field in format_fields:
                key = (False, field.id)
                if key not in existing_combinations:
                    SettingLine.create({
                        'format_setting_id': self.id,
                        'format_field_id': field.id,
                    })
                    lines_created += 1

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Columnas cargadas'),
                'message': _('%d columnas agregadas a la configuracion.') % lines_created,
                'sticky': False,
                'type': 'success' if lines_created > 0 else 'warning',
            }
        }

    def action_view_concept_accounts(self):
        """Abre el informe de conceptos y cuentas asociadas pre-filtrado"""
        self.ensure_one()
        wizard = self.env['l10n_co.exogenous_concept_accounts_report_wizard'].create({
            'format_id': self.format_id.id,
            'company_id': self.company_id.id,
        })
        wizard._generate_report_lines()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Conceptos y cuentas - %s') % self.format_id.code,
            'res_model': 'l10n_co.exogenous_concept_accounts_report_wizard',
            'view_mode': 'form',
            'res_id': wizard.id,
            'target': 'new',
        }

    def _get_xsd_path(self):
        """Obtiene la ruta al archivo XSD del formato"""
        module_path = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        xsd_file = os.path.join(module_path, 'xsd', f'{self.format_id.code}.xsd')
        if os.path.exists(xsd_file):
            return xsd_file
        return None

    def _validate_xml_with_xsd(self, xml_content):
        """Valida el XML generado contra el esquema XSD

        Prioridad de búsqueda de XSD:
        1. XSD binario subido en el formato (xsd_file)
        2. Archivo XSD en carpeta local (xsd/{codigo}.xsd)
        """
        xsd_source = None
        xsd_doc = None

        # Prioridad 1: XSD binario subido al formato
        if self.format_id.xsd_file:
            try:
                xsd_content = base64.b64decode(self.format_id.xsd_file)
                xsd_doc = etree.parse(io.BytesIO(xsd_content))
                xsd_source = 'binario'
                _logger.info("Usando XSD binario para formato %s", self.format_id.code)
            except Exception as e:
                _logger.warning("Error leyendo XSD binario: %s", str(e))

        # Prioridad 2: Archivo XSD en ruta local
        if xsd_doc is None:
            xsd_path = self._get_xsd_path()
            if xsd_path:
                try:
                    with open(xsd_path, 'rb') as xsd_file:
                        xsd_doc = etree.parse(xsd_file)
                        xsd_source = 'archivo'
                        _logger.info("Usando XSD de archivo: %s", xsd_path)
                except Exception as e:
                    _logger.warning("Error leyendo XSD de archivo: %s", str(e))

        if xsd_doc is None:
            _logger.warning("No se encontró XSD para el formato %s", self.format_id.code)
            return True, []

        try:
            xsd_schema = etree.XMLSchema(xsd_doc)
            xml_doc = etree.fromstring(xml_content)
            is_valid = xsd_schema.validate(xml_doc)

            errors = []
            if not is_valid:
                for error in xsd_schema.error_log:
                    errors.append(f"Línea {error.line}: {error.message}")

            return is_valid, errors
        except Exception as e:
            _logger.error("Error validando XML con XSD %s: %s", xsd_source, str(e))
            return False, [str(e)]

    def _build_xml_header(self, total_value, record_count):
        """Construye el elemento de cabecera del XML"""
        cab = etree.Element('Cab')

        year = self.date_end.year if self.date_end else date.today().year - 1
        etree.SubElement(cab, 'Ano').text = str(year)

        concept_code = '01'
        if self.format_setting_line_ids and self.format_id.apply_concepts:
            cl_concept_code = self.format_setting_line_ids[0].concept_id.code or '01'
            cl_concept_code = ''.join(ch for ch in cl_concept_code if ch.isdigit()) or '01'
            if len(cl_concept_code) > 2:
                cl_concept_code = cl_concept_code[-2:]
            concept_code = cl_concept_code
        etree.SubElement(cab, 'CodCpt').text = str(concept_code).zfill(2)

        etree.SubElement(cab, 'Formato').text = self.format_id.code
        etree.SubElement(cab, 'Version').text = str(self.format_id.xsd_version or 10)
        etree.SubElement(cab, 'NumEnvio').text = str(self.send_number)
        etree.SubElement(cab, 'FecEnvio').text = datetime.now().strftime('%Y-%m-%dT%H:%M:%S')
        etree.SubElement(cab, 'FecInicial').text = self.date_start.strftime('%Y-%m-%d') if self.date_start else f'{year}-01-01'
        etree.SubElement(cab, 'FecFinal').text = self.date_end.strftime('%Y-%m-%d') if self.date_end else f'{year}-12-31'
        etree.SubElement(cab, 'ValorTotal').text = str(round(total_value, 2))
        etree.SubElement(cab, 'CantReg').text = str(record_count)

        return cab

    # --- Mapeos comunes Unicode → ISO-8859-1 (chars que no caben directamente) ---
    _ISO88591_REPLACEMENTS = {
        '‘': "'", '’': "'", '‚': "'", '‛': "'",  # comillas simples tipográficas
        '“': '"', '”': '"', '„': '"', '‟': '"',  # comillas dobles tipográficas
        '–': '-', '—': '-', '―': '-',                  # guiones largos
        '…': '...',                                              # elipsis
        ' ': ' ', ' ': ' ', ' ': ' ',                  # espacios no-rompibles
        '­': '',                                                 # soft hyphen
        '​': '', '‌': '', '‍': '', '﻿': '',       # zero-width
    }

    @classmethod
    def _sanitize_iso88591(cls, value):
        """Devuelve un str que se codifica limpio a ISO-8859-1 (cp1).

        - Normaliza Unicode (NFC).
        - Reemplaza chars típicos no-ISO-8859-1 con equivalentes ASCII.
        - Quita caracteres de control no imprimibles.
        - Para lo que aún no cabe en ISO-8859-1, intenta NFKD/ASCII; si no, '?'.
        """
        if value is None:
            return ''
        s = unicodedata.normalize('NFC', str(value))
        s = ''.join(cls._ISO88591_REPLACEMENTS.get(ch, ch) for ch in s)
        s = ''.join(' ' if ch < ' ' else ch for ch in s)
        try:
            s.encode('iso-8859-1')
            return s
        except UnicodeEncodeError:
            pass
        out = []
        for ch in s:
            try:
                ch.encode('iso-8859-1')
                out.append(ch)
            except UnicodeEncodeError:
                ascii_ch = unicodedata.normalize('NFKD', ch).encode('ascii', 'ignore').decode('ascii')
                out.append(ascii_ch or '?')
        return ''.join(out)

    def _get_format_field_overrides(self, format_fields):
        """Devuelve dict {field_id: override_record} para el formato actual."""
        if not format_fields:
            return {}
        overrides = self.env['l10n_co.exogenous_format_field_override'].sudo().search([
            ('format_id', '=', self.format_id.id),
            ('field_id', 'in', format_fields.ids),
        ])
        return {o.field_id.id: o for o in overrides}

    _COLAB_SPECIAL_ATTRS = {'tcon', 'idfi', 'tdopa', 'nidpa'}

    def _build_xml_record(self, row, format_fields, contract=None, participant=None, pct=None):
        """Construye un elemento XML para cada registro.

        Args:
            row: dict con datos del movimiento/partner (lo que devuelve el dataframe)
            format_fields: recordset de format_field activos para este formato
            contract: l10n_co.exogenous_colaboracion_contract opcional (para 5247-5252)
            participant: l10n_co.exogenous_colaboracion_participant opcional
            pct: float % de participación (si aplica) — multiplica los montos
        """
        element_name = self.format_id.xml_element_name or 'registro'
        record = etree.Element(element_name)
        uvt = self.uvt_value or 0
        overrides = self._get_format_field_overrides(format_fields)

        # Ordenar por sequence efectiva (override.sequence si existe, sino field.sequence)
        def _eff_sequence(f):
            ov = overrides.get(f.id)
            return (ov.sequence if (ov and ov.sequence) else f.sequence) or 0
        ordered_fields = sorted(format_fields, key=_eff_sequence)

        # Datos especiales de colaboración (5247-5252)
        colab_values = {}
        if contract:
            colab_values['tcon'] = contract.contract_type_id.code or ''
            colab_values['idfi'] = contract.code or ''
        if participant:
            id_type = participant.partner_id.l10n_latam_identification_type_id
            colab_values['tdopa'] = id_type.l10n_co_document_code or '' if id_type else ''
            colab_values['nidpa'] = (participant.partner_id.vat or '').replace('.', '').replace('-', '')

        for field in ordered_fields:
            attr_name = field.attribute
            if not attr_name:
                continue

            ov = overrides.get(field.id)
            eff_max_length = (ov.max_length if (ov and ov.max_length) else None) or field.max_length

            if attr_name in self._COLAB_SPECIAL_ATTRS:
                value = colab_values.get(attr_name, '')
                if value == '':
                    continue
                value = self._sanitize_iso88591(str(value))
                if eff_max_length and len(value) > eff_max_length:
                    value = value[:eff_max_length]
                record.set(attr_name, value)
                continue

            # Columna derivada: toma valor de otra columna
            if field.formula_source_field_id and field.source == 'journal_items':
                source_name = field.formula_source_field_id.name
                value = row.get(source_name, 0) or 0
                if isinstance(value, (int, float)):
                    value = field.apply_value_transformations(value, uvt_value=uvt)
            else:
                value = row.get(field.name, '')

            if pd.isna(value) or value == '':
                if field.source == 'journal_items':
                    value = '0'
                else:
                    continue

            if field.source == 'journal_items' and isinstance(value, (int, float)) and not field.formula_source_field_id:
                value = field.apply_value_transformations(value, uvt_value=uvt)

            # Aplicar % de participación a los montos cuando hay partícipe
            if pct is not None and field.source == 'journal_items' and isinstance(value, (int, float)):
                value = round(value * (pct / 100.0))

            if isinstance(value, float):
                value = str(int(value)) if value == int(value) else str(value)
            else:
                value = str(value)

            # Aplicar recorte (xsd_slice) ANTES del truncate
            if field.xsd_slice and field.xsd_slice != 'none':
                value = str(field.apply_slice(value))

            value = self._sanitize_iso88591(value)

            if eff_max_length and len(value) > eff_max_length:
                value = value[:eff_max_length]

            record.set(attr_name, value)

        return record

    def _apply_manual_column_overrides(self, account_move_lines, format_field):
        """Aplica asignaciones manuales (l10n_co_exogenous_format_field_id) a las líneas.

        Si la línea tiene asignación manual que coincide con el format_field actual,
        se respeta el porcentaje configurado. Si está asignada a otro campo, se excluye
        de este campo (para evitar doble conteo). Si tiene prorrateo (pct < 100), se
        reemplaza el balance por balance * pct / 100.

        Limitaciones actuales: este hook trata el caso primaria; la porción secundaria
        (residuo) se contabiliza cuando el generador procesa el otro format_field.
        Refactor completo de agrupación queda para siguiente iteración.
        """
        if not format_field or not account_move_lines:
            return account_move_lines

        result = []
        for row in account_move_lines:
            line_id = row.get('id')
            if not line_id:
                result.append(row)
                continue

            line = self.env['account.move.line'].browse(line_id)
            assigned = line.l10n_co_exogenous_format_field_id
            if not assigned:
                result.append(row)
                continue

            pct = line.l10n_co_exogenous_percentage or 100.0
            secondary = line.l10n_co_exogenous_secondary_field_id

            if assigned.id == format_field.id:
                if pct < 100 and 'balance' in row:
                    new_row = dict(row)
                    factor = pct / 100.0
                    for k in ('balance', 'debit', 'credit', 'tax_base_amount'):
                        if k in new_row and isinstance(new_row[k], (int, float)):
                            new_row[k] = new_row[k] * factor
                    result.append(new_row)
                else:
                    result.append(row)
            elif secondary and secondary.id == format_field.id and pct < 100:
                new_row = dict(row)
                factor = (100.0 - pct) / 100.0
                for k in ('balance', 'debit', 'credit', 'tax_base_amount'):
                    if k in new_row and isinstance(new_row[k], (int, float)):
                        new_row[k] = new_row[k] * factor
                result.append(new_row)
            # else: asignada a otro field → excluir de este campo
        return result

    def generate_and_download_xml(self):
        """Genera el archivo XML con validación XSD"""
        self.ensure_one()
        # Garantizar contexto ORM correcto: empresa principal + sucursales si aplica
        self = self.with_company(self.company_id | self.company_ids)

        if not self.format_id.xml_element_name:
            return self._show_message_error(
                f"El formato {self.format_id.code} no tiene configurado el elemento XML"
            )

        if not self.format_setting_line_ids:
            return self._show_message_error(
                f"No hay configuración de formato para: {self.format_id.code}"
            )

        format_fields = self.env['l10n_co.exogenous_format_field'].search(
            [('format_ids', 'in', self.format_id.id)], order="sequence asc")

        if not format_fields:
            return self._show_message_error(
                f"No hay campos configurados para el formato: {self.format_id.code}"
            )

        fields_contact = format_fields.filtered(lambda ff: ff.source == 'contact')
        fields_to_clear_contact = self._get_fields_to_clear_contact(fields_contact)
        fields_to_clear_company = self._get_fields_to_clear_company(fields_contact)
        unique_keys_by_format = format_fields.filtered(
            lambda ff: ff.is_unique_key).mapped('field_odoo_id').mapped('name')

        if self.format_id.apply_concepts:
            unique_keys_by_format = unique_keys_by_format + ['account_id']
        if any(sl.group_by_journal for sl in self.format_setting_line_ids):
            unique_keys_by_format = unique_keys_by_format + ['journal_id']

        concepts = self._get_accounts_by_format()
        all_data = []

        for setting_line in self.format_setting_line_ids:
            accounts_ids = self._get_accounts(setting_line).ids
            partner_ids, account_move_lines = self._get_information_by_account_move_line(accounts_ids, format_field=setting_line.format_field_id, setting_line=setting_line)
            account_move_lines = self._apply_manual_column_overrides(account_move_lines, setting_line.format_field_id)
            account_move_lines, partner_ids = self._resolve_partner_grouping(account_move_lines, partner_ids)

            if not account_move_lines or not partner_ids:
                continue

            partners, other_info_contact = self._get_information_partner(partner_ids, fields_contact)
            if isinstance(partners, dict) and partners.get('tag', False):
                return partners

            df_type_documents = pd.DataFrame(self._get_type_documents_by_format())
            df_account_move_lines = pd.DataFrame(account_move_lines)
            df_partners = pd.DataFrame(partners)
            df_partners = df_partners.fillna('').replace(False, '')
            df_other_info = pd.DataFrame([other_info_contact])

            df_partners = self._normalice_data_dataframe_partner(
                self._get_field_name_second_field_many2one(fields_contact),
                df_partners, df_type_documents, df_other_info,
                fields_to_clear_contact, fields_to_clear_company
            )

            df_account_move_lines = self._normalice_and_merge_data_dataframe_account_move_line(
                df_account_move_lines, df_partners)

            format_field_accumulated = self._format_field_by_accumulated(setting_line)

            def get_column_value(row, key):
                """Calcula el valor de la columna según la fórmula configurada"""
                column_to_get = format_field_accumulated.get(key, None)
                if isinstance(column_to_get, dict):
                    account_id = row['account_id']
                    formula = column_to_get.get(account_id, None)
                    if formula:
                        return self._compute_formula(row, formula)
                return 0.0

            columns_ordered = self._get_columns_ordered(format_fields)
            columns_ordered = columns_ordered + [setting_line.format_field_id.name]
            clave_diccionario, _valor = next(iter(format_field_accumulated.items()))

            df_account_move_lines[clave_diccionario] = df_account_move_lines.apply(
                lambda row: get_column_value(row, clave_diccionario), axis=1)

            default_concept_code = (
                setting_line.concept_id.code
                if (self.format_id.apply_concepts and setting_line.concept_id)
                else ''
            )
            df_account_move_lines['account_id'] = df_account_move_lines['account_id'].apply(
                lambda x: concepts.get(x[0], default_concept_code))

            operations = {
                column: 'first' if column not in self._get_fields_format_id(setting_line) else 'sum'
                for column in columns_ordered
            }

            df_copy = df_account_move_lines.copy(deep=True)
            filtered_grouped = df_copy.groupby(unique_keys_by_format)
            result = filtered_grouped.agg(operations).round(0)
            result = result.rename(columns=self._get_fields_odoo_and_format(fields_contact))

            all_data.append(result)

        if not all_data:
            return self._show_message_error("No se encontró información con estos parámetros")

        final_df = pd.concat(all_data, ignore_index=True)

        # Merge final: cada setting_line genera una fila aparte, pero un mismo
        # tercero/concepto debe aparecer una sola vez con todos los campos llenos.
        journal_field_names = format_fields.filtered(
            lambda f: f.source == 'journal_items').mapped('name')
        unique_field_names = format_fields.filtered(
            lambda f: f.is_unique_key).mapped('name')
        if self.format_id.apply_concepts:
            cpt_field = self._get_field_concept()
            if cpt_field and cpt_field.name not in unique_field_names:
                unique_field_names = unique_field_names + [cpt_field.name]
        group_keys = [c for c in unique_field_names if c in final_df.columns]
        if group_keys:
            agg = {}
            for c in final_df.columns:
                if c in group_keys:
                    continue
                if c in journal_field_names:
                    agg[c] = 'sum'
                else:
                    agg[c] = 'first'
            if agg:
                fillna_cols = {c: 0 for c in journal_field_names if c in final_df.columns}
                if fillna_cols:
                    final_df = final_df.fillna(value=fillna_cols)
                final_df = final_df.groupby(group_keys, as_index=False, dropna=False).agg(agg)

        # Columnas derivadas y transformaciones de porcentaje/limite
        uvt = self.uvt_value or 0
        journal_fields = format_fields.filtered(lambda f: f.source == 'journal_items')
        for field in journal_fields:
            if field.name not in final_df.columns:
                continue
            if field.formula_source_field_id:
                source_col = field.formula_source_field_id.name
                if source_col in final_df.columns:
                    final_df[field.name] = final_df[source_col].apply(
                        lambda v: field.apply_value_transformations(v, uvt_value=uvt)
                        if isinstance(v, (int, float)) else v)
            elif field.formula_percentage != 100.0 or field.formula_limit_type != 'none':
                final_df[field.name] = final_df[field.name].apply(
                    lambda v: field.apply_value_transformations(v, uvt_value=uvt)
                    if isinstance(v, (int, float)) else v)

        root_element = self.format_id.xml_root_element or 'mas'
        root = etree.Element(root_element)

        total_fields = format_fields.filtered(lambda f: f.source == 'journal_items')
        total_value = 0
        if total_fields:
            first_total_field = total_fields[0].name
            if first_total_field in final_df.columns:
                total_value = final_df[first_total_field].sum()

        contract = self.colaboracion_contract_id
        participants = contract.participant_ids if contract else None

        if contract and participants:
            record_count = len(final_df) * len(participants)
        else:
            record_count = len(final_df)

        cab = self._build_xml_header(total_value, record_count)
        root.append(cab)

        for _idx, row in final_df.iterrows():
            row_dict = row.to_dict()
            if contract and participants:
                for p in participants:
                    record = self._build_xml_record(
                        row_dict, format_fields,
                        contract=contract, participant=p, pct=p.participation_pct)
                    root.append(record)
            else:
                record = self._build_xml_record(row_dict, format_fields)
                root.append(record)

        xml_content = etree.tostring(
            root, pretty_print=True, xml_declaration=True, encoding='ISO-8859-1')

        is_valid, errors = self._validate_xml_with_xsd(xml_content)
        if not is_valid:
            error_msg = "Errores de validación XSD:\n" + "\n".join(errors[:10])
            _logger.warning(error_msg)

        file_name = f"dmuisca_{self.format_id.code}_{datetime.now().strftime('%Y%m%d%H%M%S')}"
        self.xml_file = base64.b64encode(xml_content)
        self.xml_file_name = f"{file_name}.xml"
        self.send_number += 1

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/?model=%s&id=%d&field=xml_file&filename_field=xml_file_name&download=true' % (
                self._name, self.id),
            'target': 'self',
        }

    def generate_and_download_flat_file(self):
        """Genera archivo plano (texto delimitado) para exógena distrital."""
        self.ensure_one()
        # Garantizar contexto ORM correcto: empresa principal + sucursales si aplica
        self = self.with_company(self.company_id | self.company_ids)
        district = self.format_id.district_id
        if not district:
            return self._show_message_error(
                "Este formato no tiene distrito/municipio configurado. "
                "Solo los formatos distritales generan archivo plano.")

        if not self.format_setting_line_ids:
            return self._show_message_error(
                f"No hay configuración de formato para: {self.format_id.code}")

        format_fields = self.env['l10n_co.exogenous_format_field'].search(
            [('format_ids', 'in', self.format_id.id)], order="sequence asc")
        if not format_fields:
            return self._show_message_error(
                f"No hay campos configurados para el formato: {self.format_id.code}")

        # --- Reutilizar pipeline de datos del Excel ---
        fields_contact = format_fields.filtered(lambda ff: ff.source == 'contact')
        fields_to_clear_contact = self._get_fields_to_clear_contact(fields_contact)
        fields_to_clear_company = self._get_fields_to_clear_company(fields_contact)
        unique_keys_by_format = format_fields.filtered(
            lambda ff: ff.is_unique_key).mapped('field_odoo_id').mapped('name')
        unique_keys_by_format_field = format_fields.filtered(
            lambda ff: ff.is_unique_key).mapped('name')

        if self.format_id.apply_concepts:
            unique_keys_by_format = unique_keys_by_format + ['account_id']

        original_dataframe = self._create_original_dataframe(format_fields)
        columnas_originales = original_dataframe.columns
        row_smaller_amount = self._get_fields_fill_smaller_amount()
        concepts = self._get_accounts_by_format()
        get_information = list()

        for setting_line in self.format_setting_line_ids:
            accounts_ids = self._get_accounts(setting_line).ids
            partner_ids, account_move_lines = self._get_information_by_account_move_line(
                accounts_ids, format_field=setting_line.format_field_id, setting_line=setting_line)
            account_move_lines, partner_ids = self._resolve_partner_grouping(account_move_lines, partner_ids)

            if not account_move_lines or not partner_ids:
                get_information.append(False)
                continue
            else:
                get_information.append(True)

            partners, other_info_contact = self._get_information_partner(partner_ids, fields_contact)
            if isinstance(partners, dict) and partners.get('tag', False):
                return partners

            df_type_documents = pd.DataFrame(self._get_type_documents_by_format())
            df_account_move_lines = pd.DataFrame(account_move_lines)
            df_partners = pd.DataFrame(partners)
            df_partners = df_partners.fillna('')
            df_partners = df_partners.replace(False, '')
            df_other_info = pd.DataFrame([other_info_contact])

            df_partners = self._normalice_data_dataframe_partner(
                self._get_field_name_second_field_many2one(fields_contact),
                df_partners, df_type_documents, df_other_info,
                fields_to_clear_contact, fields_to_clear_company)

            df_account_move_lines = self._normalice_and_merge_data_dataframe_account_move_line(
                df_account_move_lines, df_partners)

            format_field_accumulated = self._format_field_by_accumulated(setting_line)

            def get_column_value(row, key):
                column_to_get = format_field_accumulated.get(key, None)
                if isinstance(column_to_get, dict):
                    account_id = row['account_id']
                    formula = column_to_get.get(account_id, None)
                    if formula:
                        return self._compute_formula(row, formula)
                return 0.0

            columns_ordered = self._get_columns_ordered(format_fields)
            columns_ordered = columns_ordered + [setting_line.format_field_id.name]
            clave_diccionario, valor_diccionario = next(iter(format_field_accumulated.items()))
            df_account_move_lines[clave_diccionario] = df_account_move_lines.apply(
                lambda row: get_column_value(row, clave_diccionario), axis=1)
            default_concept_code = (
                setting_line.concept_id.code
                if (self.format_id.apply_concepts and setting_line.concept_id)
                else ''
            )
            df_account_move_lines['account_id'] = df_account_move_lines['account_id'].apply(
                lambda x: concepts.get(x[0], default_concept_code))

            operations = {column: 'first' if column not in self._get_fields_format_id(setting_line) else 'sum'
                          for column in columns_ordered}
            df_copy = df_account_move_lines.copy(deep=True)
            filtered_grouped = df_copy.groupby(unique_keys_by_format)
            result = filtered_grouped.agg(operations).round(0)
            result = result.rename(columns=self._get_fields_odoo_and_format(fields_contact))

            if self.format_id.applying_smaller_amounts:
                # Compare on absolute value: see explanation in generate_and_download_report
                col = setting_line.format_field_id.name
                _abs = result[col].abs()
                smaller_amounts = result.loc[_abs < self.format_id.smaller_ammounts, col].sum()
                result = result.loc[_abs >= self.format_id.smaller_ammounts]
                if self.format_id.apply_concepts and smaller_amounts > 0:
                    row_sm = self._get_fields_fill_smaller_amount()
                    row_sm.update({setting_line.format_field_id.name: smaller_amounts})
                    row_sm.update({self._get_field_concept().name: setting_line.concept_id.code})
                    result = self._generate_row_by_smaller_amount(row_sm, result)
                elif not self.format_id.apply_concepts and smaller_amounts > 0:
                    row_smaller_amount.update({setting_line.format_field_id.name: smaller_amounts})

            result = result.reset_index(drop=True)
            duplicados = original_dataframe.merge(result, on=unique_keys_by_format_field, how='inner')
            filas_duplicadas = result[result.index.isin(duplicados.index)]
            for index, fila_duplicada in filas_duplicadas.iterrows():
                filtro = (original_dataframe[unique_keys_by_format_field] == fila_duplicada[unique_keys_by_format_field]).all(axis=1)
                original_dataframe.loc[filtro, fila_duplicada.index] = fila_duplicada.values
            filas_nuevas = result[~result.index.isin(duplicados.index)]
            original_dataframe = pd.concat([original_dataframe, filas_nuevas], ignore_index=True)

        if not any(get_information):
            return self._show_message_error("No se encontró información con estos parámetros")

        if not self.format_id.apply_concepts and self.format_id.applying_smaller_amounts:
            original_dataframe = self._generate_row_by_smaller_amount(row_smaller_amount, original_dataframe)

        original_dataframe = original_dataframe.reindex(columns=columnas_originales)

        # DIAN espera ceros en columnas de monto, no celdas vacías.
        amount_fields = format_fields.filtered(lambda f: f.source == 'journal_items').mapped('name')
        for col in amount_fields:
            if col in original_dataframe.columns:
                original_dataframe[col] = original_dataframe[col].fillna(0)

        # Aplicar recorte (xsd_slice) ANTES de sanitizar/truncar
        for ff in format_fields:
            if ff.xsd_slice and ff.xsd_slice != 'none' and ff.name in original_dataframe.columns:
                original_dataframe[ff.name] = original_dataframe[ff.name].apply(
                    lambda v, f=ff: f.apply_slice(v))

        # --- Sanitización ---
        if district.sanitize_text:
            from odoo.addons.l10n_co_exogenous_information_reporting.tools.sanitize import sanitize_dataframe
            original_dataframe = sanitize_dataframe(
                original_dataframe,
                sanitize=True,
                max_length=district.max_field_length)

        # --- Generar archivo plano ---
        delimiter = district._get_delimiter_char()
        encoding = district.text_encoding or 'utf-8'
        extension = district.file_extension or 'txt'

        output = io.StringIO()
        column_names = format_fields.mapped('name')

        if district.include_header:
            output.write(delimiter.join(column_names) + '\n')

        for _idx, row in original_dataframe.iterrows():
            values = []
            for col in column_names:
                val = row.get(col, '')
                if val is None or (isinstance(val, float) and pd.isna(val)):
                    val = ''
                else:
                    val = str(val).strip()
                values.append(val)
            output.write(delimiter.join(values) + '\n')

        content = output.getvalue()
        file_bytes = content.encode(encoding, errors='replace')
        year = self.date_start.year if self.date_start else 'YYYY'
        filename = f"exogena_distrital_{district.code}_{self.format_id.code}_{year}.{extension}"

        self.write({
            'flat_file': base64.b64encode(file_bytes),
            'flat_file_name': filename,
        })

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/?model=%s&id=%d&field=flat_file&filename_field=flat_file_name&download=true' % (
                self._name, self.id),
            'target': 'self',
        }

    def _get_xsd_document(self):
        """Obtiene el documento XSD parseado (binario o archivo)"""
        # Prioridad 1: XSD binario
        if self.format_id.xsd_file:
            try:
                xsd_content = base64.b64decode(self.format_id.xsd_file)
                return etree.parse(io.BytesIO(xsd_content))
            except Exception as e:
                _logger.warning("Error leyendo XSD binario: %s", str(e))

        # Prioridad 2: Archivo local
        xsd_path = self._get_xsd_path()
        if xsd_path:
            try:
                with open(xsd_path, 'rb') as xsd_file:
                    return etree.parse(xsd_file)
            except Exception as e:
                _logger.warning("Error leyendo XSD de archivo: %s", str(e))

        return None

    def action_audit_configuration(self):
        """Abre el wizard de auditoría completa del formato"""
        self.ensure_one()

        if not self.format_id:
            return self._show_message_error("Debe seleccionar un formato primero")

        wizard = self.env['l10n_co.exogenous_audit_wizard'].create({
            'format_setting_id': self.id,
        })
        return wizard.action_run_audit()

    def action_preview_columns(self):
        """Muestra preview de columnas del formato con restricciones XSD y cuentas asignadas"""
        self.ensure_one()

        if not self.format_id:
            return self._show_message_error("Debe seleccionar un formato primero")

        format_fields = self.env['l10n_co.exogenous_format_field'].search([
            ('format_ids', 'in', self.format_id.id),
        ], order="sequence asc")

        if not format_fields:
            return self._show_message_error(
                f"No hay campos configurados para el formato {self.format_id.code}"
            )

        # Construir mapa setting_line por field_id
        lines_by_field = {}
        for line in self.format_setting_line_ids:
            fid = line.format_field_id.id
            if fid not in lines_by_field:
                lines_by_field[fid] = []
            lines_by_field[fid].append(line)

        # Construir HTML del preview
        rows_html = ""
        for field in format_fields:
            xsd_info = self._get_xsd_restriction_summary(field)
            lines = lines_by_field.get(field.id, [])
            accounts_parts = []
            for line in lines:
                accts = self._get_accounts(line)
                if accts:
                    codes = ', '.join(accts.mapped('code')[:10])
                    if len(accts) > 10:
                        codes += f" ... (+{len(accts) - 10})"
                    label = line.concept_id.code if line.concept_id else ''
                    if label:
                        accounts_parts.append(f"<b>{label}</b>: {codes}")
                    else:
                        accounts_parts.append(codes)

            accounts_html = '<br/>'.join(accounts_parts) if accounts_parts else '<i style="color:#999">Sin cuentas</i>'
            req = '&#10003;' if field.xsd_required else ''
            rows_html += f"""
            <tr>
                <td style="padding:4px 8px">{field.sequence}</td>
                <td style="padding:4px 8px"><b>{field.attribute or ''}</b></td>
                <td style="padding:4px 8px">{field.name or ''}</td>
                <td style="padding:4px 8px">{field.source or ''}</td>
                <td style="padding:4px 8px;text-align:center">{req}</td>
                <td style="padding:4px 8px;font-size:11px">{xsd_info}</td>
                <td style="padding:4px 8px;font-size:11px">{accounts_html}</td>
            </tr>"""

        html_content = f"""
        <div style="max-height:500px;overflow:auto">
        <table class="table table-sm table-striped" style="font-size:12px">
            <thead style="position:sticky;top:0;background:#dee2e6">
                <tr>
                    <th style="padding:4px 8px">Seq</th>
                    <th style="padding:4px 8px">Atributo</th>
                    <th style="padding:4px 8px">Nombre</th>
                    <th style="padding:4px 8px">Fuente</th>
                    <th style="padding:4px 8px">Req</th>
                    <th style="padding:4px 8px">Restricciones XSD</th>
                    <th style="padding:4px 8px">Cuentas asignadas</th>
                </tr>
            </thead>
            <tbody>{rows_html}</tbody>
        </table>
        </div>
        """

        # Usar wizard transient para mostrar el HTML
        wizard = self.env['l10n_co.exogenous_column_preview_wizard'].create({
            'format_setting_id': self.id,
            'preview_html': html_content,
        })
        return {
            'type': 'ir.actions.act_window',
            'name': _('Preview columnas - %s') % self.format_id.code,
            'res_model': 'l10n_co.exogenous_column_preview_wizard',
            'view_mode': 'form',
            'res_id': wizard.id,
            'target': 'new',
        }

    def action_detail_audit(self):
        """Abre el auxiliar detallado con lineas de movimiento por concepto"""
        self.ensure_one()

        if not self.format_id:
            return self._show_message_error("Debe seleccionar un formato primero")

        if not self.format_setting_line_ids:
            return self._show_message_error("No hay líneas de configuración. Configure las columnas primero.")

        if self.is_it_with_date_range and (not self.date_start or not self.date_end):
            return self._show_message_error("Configure las fechas del período antes de generar el auxiliar")

        wizard = self.env['l10n_co.exogenous_detail_audit_wizard'].create({
            'format_setting_id': self.id,
        })
        return wizard.action_generate_detail()

    def _build_partner_audit_domain(self):
        """Construye el dominio combinado de todas las setting_lines para el auxiliar por tercero.
        Une las cuentas de todas las lineas y aplica filtros comunes."""
        self.ensure_one()

        # Recopilar todas las cuentas de todas las setting lines
        all_account_ids = set()
        for setting_line in self.format_setting_line_ids:
            accounts = self._get_accounts(setting_line)
            if accounts:
                all_account_ids.update(accounts.ids)

        if not all_account_ids:
            return []

        rc = self._get_reporting_company_ids()
        domain = [
            ('parent_state', '=', 'posted'),
            ('company_id', 'in' if len(rc) > 1 else '=', rc if len(rc) > 1 else rc[0]),
            ('account_id', 'in', list(all_account_ids)),
            ('partner_id', '!=', False),
        ]

        # Fechas: usar rango mas amplio (sin corte por campo para la vista dinamica)
        if not self.is_it_with_date_range:
            domain.append(('date', '<=', self._get_last_day_pass_year()))
        else:
            if self.date_start:
                domain.append(('date', '>=', str(self.date_start)))
            if self.date_end:
                domain.append(('date', '<=', str(self.date_end)))

        # Exclusiones estandar
        if self.journal_ids:
            domain.append(('journal_id', 'not in', self.journal_ids.ids))
        if self.tax_ids:
            domain.append(('tax_line_id', 'not in', self.tax_ids.ids))
        if self.partner_ids:
            domain.append(('partner_id', 'not in', self.partner_ids.ids))

        # Dominio personalizado global
        if self.use_custom_domain and self.custom_domain:
            try:
                custom = self._parse_custom_domain(self.custom_domain)
                domain += custom
            except (ValidationError, Exception):
                pass

        return domain

    def action_partner_audit(self):
        """Abre vista dinamica de account.move.line agrupada por tercero"""
        self.ensure_one()

        if not self.format_id:
            return self._show_message_error("Debe seleccionar un formato primero")

        if not self.format_setting_line_ids:
            return self._show_message_error("No hay lineas de configuracion. Configure las columnas primero.")

        if self.is_it_with_date_range and (not self.date_start or not self.date_end):
            return self._show_message_error("Configure las fechas del periodo antes de generar el auxiliar")

        domain = self._build_partner_audit_domain()
        if not domain:
            return self._show_message_error("No se encontraron cuentas configuradas en las lineas")

        search_view = self.env.ref(
            'l10n_co_exogenous_information_reporting.view_account_move_line_exogenous_partner_search',
            raise_if_not_found=False)

        return {
            'type': 'ir.actions.act_window',
            'name': _('Aux. Tercero - %s') % self.format_id.code,
            'res_model': 'account.move.line',
            'view_mode': 'list,pivot,graph',
            'domain': domain,
            'search_view_id': search_view.id if search_view else False,
            'context': {
                'expand': False,
                'search_default_group_partner': 1,
            },
            'target': 'current',
        }

    def action_partner_audit_excel(self):
        """Abre el wizard de exportacion Excel del auxiliar por tercero"""
        self.ensure_one()

        if not self.format_id:
            return self._show_message_error("Debe seleccionar un formato primero")

        if not self.format_setting_line_ids:
            return self._show_message_error("No hay lineas de configuracion. Configure las columnas primero.")

        if self.is_it_with_date_range and (not self.date_start or not self.date_end):
            return self._show_message_error("Configure las fechas del periodo antes de generar el auxiliar")

        wizard = self.env['l10n_co.exogenous_partner_audit_wizard'].create({
            'format_setting_id': self.id,
        })
        return wizard.action_generate_partner_detail()
