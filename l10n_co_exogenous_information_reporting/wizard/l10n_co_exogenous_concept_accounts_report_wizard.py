# -*- coding: utf-8 -*-
import base64
import io
import logging
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from odoo import api, fields, models, _, Command
from odoo.exceptions import UserError
from odoo.addons.l10n_co_exogenous_information_reporting.models.l10n_co_exogenous_format_field_account import (
    NATURE_ACCOUNT_SELECTION,
)

_logger = logging.getLogger(__name__)


class L10nCoExogenousConceptAccountsReportWizard(models.TransientModel):
    """Wizard para generar el informe de conceptos y cuentas asociadas"""
    _name = 'l10n_co.exogenous_concept_accounts_report_wizard'
    _description = 'Informe de conceptos y cuentas asociadas'

    format_id = fields.Many2one(
        'l10n_co.exogenous_format',
        string='Reporte',
        help='Seleccione un formato específico o deje vacío para mostrar todos'
    )
    company_id = fields.Many2one(
        'res.company',
        string='Compañía',
        default=lambda self: self.env.company,
        required=True
    )
    format_field_filter_id = fields.Many2one(
        'l10n_co.exogenous_format_field',
        string='Columna',
        help='Filtrar por columna específica del reporte'
    )
    detail_columns = fields.Boolean(
        string='Detalle por columnas',
        default=False,
        help='Cuando está marcado, muestra la información organizada por columnas en lugar de conceptos'
    )
    line_ids = fields.One2many(
        'l10n_co.exogenous_concept_accounts_report_wizard.line',
        'wizard_id',
        string='Resultados'
    )
    binary_file = fields.Binary(string='Archivo Excel')
    binary_file_name = fields.Char(string='Nombre del archivo')

    @api.onchange('format_id', 'format_field_filter_id', 'detail_columns')
    def _onchange_format(self):
        """Genera las líneas del informe cuando cambia el formato"""
        self.line_ids = [Command.clear()]
        self._generate_report_lines()

    def action_generate_report(self):
        """Acción para generar el reporte"""
        self.ensure_one()
        self._generate_report_lines()
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
            'context': self.env.context,
        }

    def _generate_report_lines(self):
        """Genera las líneas del informe según los filtros

        Informe de "Conceptos y cuentas asociadas":
        - Muestra cuentas agrupadas por concepto o por columna
        - Incluye información de columna y naturaleza
        """
        lines = []

        Concept = self.env['l10n_co.exogenous_concept']
        domain = [*Concept._check_company_domain(self.company_id)]
        if self.format_id:
            domain.append(('format_id', '=', self.format_id.id))

        concepts = Concept.search(domain, order='format_id, code')

        for concept in concepts:
            format_name = concept.format_id.display_name if concept.format_id else ''
            concept_display = f"[{concept.code}] {concept.name}" if concept.name else f"[{concept.code}]"

            for field_account in concept.field_account_ids:
                if self.format_field_filter_id and field_account.format_field_id != self.format_field_filter_id:
                    continue

                accounts = field_account.account_ids
                if not accounts:
                    continue

                column_name = field_account.format_field_id.name if field_account.format_field_id else ''
                nature = field_account.nature_account or 'db_cr'

                if self.detail_columns:
                    header_display = f"Columna: {column_name} - Concepto: {concept_display}"
                else:
                    header_display = f"{format_name} [ {concept_display} ]"

                for account in accounts:
                    lines.append(Command.create({
                        'format_id': concept.format_id.id if concept.format_id else False,
                        'concept_id': concept.id,
                        'format_field_id': field_account.format_field_id.id if field_account.format_field_id else False,
                        'account_id': account.id,
                        'account_code': account.code,
                        'account_name': account.name,
                        'header_display': header_display,
                        'concept_code': concept.code,
                        'concept_name': concept.name or '',
                        'column_name': column_name,
                        'nature_account': nature,
                    }))

        self.line_ids = lines

    def action_export_excel(self):
        """Exporta el informe a Excel"""
        self.ensure_one()

        if not self.line_ids:
            self._generate_report_lines()

        if not self.line_ids:
            raise UserError(_('No hay datos para exportar.'))

        wb = Workbook()
        ws = wb.active
        ws.title = "Conceptos y Cuentas"

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        group_font = Font(bold=True, color='FFFFFF')
        group_fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')

        nature_labels = {
            'debit': 'DB', 'credit': 'CR',
            'db_cr': 'DB-CR', 'cr_db': 'CR-DB',
            'base_calc_db': 'Base DB', 'base_calc_cr': 'Base CR',
            'base_calc_db_cr': 'Base DB-CR', 'base_calc_cr_db': 'Base CR-DB',
            'tax_base_amount': 'Base almacenada', 'balance': 'Saldo',
        }

        headers = ['No. Cuenta Contable', 'Nombre Cuenta Contable', 'Columna', 'Naturaleza']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 50
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 15

        row = 2
        current_header = None

        for line in self.line_ids.sorted(key=lambda l: (l.header_display or '', l.account_code or '')):
            if line.header_display != current_header:
                current_header = line.header_display
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                cell = ws.cell(row=row, column=1, value=current_header)
                cell.font = group_font
                cell.fill = group_fill
                cell.alignment = Alignment(horizontal='left', vertical='center')
                cell.border = thin_border
                row += 1

            ws.cell(row=row, column=1, value=line.account_code).border = thin_border
            ws.cell(row=row, column=2, value=line.account_name).border = thin_border
            ws.cell(row=row, column=3, value=line.column_name or '').border = thin_border
            ws.cell(row=row, column=4, value=nature_labels.get(line.nature_account, '')).border = thin_border
            row += 1

        output = io.BytesIO()
        wb.save(output)

        format_code = self.format_id.code if self.format_id else 'todos'
        file_name = f"conceptos_cuentas_{format_code}.xlsx"

        self.binary_file = base64.b64encode(output.getvalue())
        self.binary_file_name = file_name

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{self._name}/{self.id}/binary_file/{file_name}?download=true',
            'target': 'self',
        }


class L10nCoExogenousConceptAccountsReportWizardLine(models.TransientModel):
    """Línea del informe de conceptos y cuentas"""
    _name = 'l10n_co.exogenous_concept_accounts_report_wizard.line'
    _description = 'Línea del informe de conceptos y cuentas'
    _order = 'header_display, account_code'

    wizard_id = fields.Many2one(
        'l10n_co.exogenous_concept_accounts_report_wizard',
        string='Asistente',
        required=True,
        ondelete='cascade'
    )
    format_id = fields.Many2one(
        'l10n_co.exogenous_format',
        string='Formato'
    )
    concept_id = fields.Many2one(
        'l10n_co.exogenous_concept',
        string='Concepto'
    )
    format_field_id = fields.Many2one(
        'l10n_co.exogenous_format_field',
        string='Columna'
    )
    account_id = fields.Many2one(
        'account.account',
        string='Cuenta'
    )
    account_code = fields.Char(
        string='Código de cuenta'
    )
    account_name = fields.Char(
        string='Nombre de cuenta'
    )
    header_display = fields.Char(
        string='Encabezado de grupo'
    )
    concept_code = fields.Char(
        string='Código de concepto'
    )
    concept_name = fields.Char(
        string='Nombre de concepto'
    )
    column_name = fields.Char(
        string='Nombre columna'
    )
    nature_account = fields.Selection(
        selection=NATURE_ACCOUNT_SELECTION,
        string='Formula'
    )
