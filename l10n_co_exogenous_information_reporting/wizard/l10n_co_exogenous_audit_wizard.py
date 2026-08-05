# -*- coding: utf-8 -*-
import base64
import io
import logging
from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from odoo import api, fields, models, _, Command
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class L10nCoExogenousAuditWizard(models.TransientModel):
    """Wizard de auditoría para validar configuración de formatos exógenos"""
    _name = 'l10n_co.exogenous_audit_wizard'
    _description = 'Auditoría de configuración de formato exógeno'

    format_setting_id = fields.Many2one(
        'l10n_co.exogenous_format_setting',
        string='Configuración',
        required=True
    )
    format_id = fields.Many2one(
        'l10n_co.exogenous_format',
        string='Formato',
        related='format_setting_id.format_id',
        readonly=True
    )
    company_id = fields.Many2one(
        'res.company',
        string='Compañía',
        related='format_setting_id.company_id',
        readonly=True
    )

    column_line_ids = fields.One2many(
        'l10n_co.exogenous_audit_wizard.column',
        'wizard_id',
        string='Columnas del reporte'
    )
    issue_line_ids = fields.One2many(
        'l10n_co.exogenous_audit_wizard.issue',
        'wizard_id',
        string='Problemas detectados'
    )

    total_columns = fields.Integer(string='Total columnas', readonly=True)
    total_fixed = fields.Integer(string='Columnas fijo', readonly=True)
    total_value = fields.Integer(string='Columnas valor', readonly=True)
    total_issues = fields.Integer(string='Problemas', readonly=True)
    total_warnings = fields.Integer(string='Advertencias', readonly=True)
    total_concepts = fields.Integer(string='Conceptos', readonly=True)
    total_accounts = fields.Integer(string='Cuentas configuradas', readonly=True)
    audit_result = fields.Selection([
        ('ok', 'Configuración válida'),
        ('warning', 'Con advertencias'),
        ('error', 'Con errores'),
    ], string='Resultado', readonly=True)

    binary_file = fields.Binary(string='Archivo Excel')
    binary_file_name = fields.Char(string='Nombre del archivo')

    def action_run_audit(self):
        """Ejecuta la auditoría completa"""
        self.ensure_one()
        self.column_line_ids = [Command.clear()]
        self.issue_line_ids = [Command.clear()]

        columns = []
        issues = []

        format_obj = self.format_id
        setting = self.format_setting_id

        # Obtener campos del formato
        format_fields = self.env['l10n_co.exogenous_format_field'].search([
            ('format_ids', 'in', format_obj.id)
        ], order='sequence asc')

        # Obtener atributos XSD
        xsd_attrs = self._get_xsd_attributes()

        # Generar líneas de columnas
        for field in format_fields:
            source_label = 'Fijo' if field.source == 'contact' else 'Valor'
            xsd_status = 'ok'
            xsd_info = ''

            if xsd_attrs:
                attr_xsd = xsd_attrs.get(field.attribute)
                if attr_xsd:
                    xsd_status = 'ok'
                    xsd_info = 'Requerido' if attr_xsd.get('required') else 'Opcional'
                elif field.attribute:
                    xsd_status = 'extra'
                    xsd_info = 'No está en XSD'
                else:
                    xsd_status = 'missing_attr'
                    xsd_info = 'Sin atributo'

            # Contar cuentas para campos de valor
            account_count = 0
            nature = ''
            if field.source == 'journal_items':
                if format_obj.apply_concepts:
                    for sl in setting.format_setting_line_ids:
                        if sl.format_field_id.id == field.id:
                            concept = sl.concept_id
                            if concept:
                                fa = concept.field_account_ids.filtered(
                                    lambda x: x.format_field_id.id == field.id
                                )
                                if fa:
                                    account_count += len(fa[0].account_ids)
                                    nature = fa[0].nature_account or ''
                else:
                    account_count = field.account_count
                    nature = field.nature_account or ''

            formula_labels = {
                'debit': 'DB', 'credit': 'CR', 'db_cr': 'DB-CR', 'cr_db': 'CR-DB',
                'base_calc_db': 'Base DB', 'base_calc_cr': 'Base CR',
                'base_calc_db_cr': 'Base DB-CR', 'base_calc_cr_db': 'Base CR-DB',
                'tax_base_amount': 'Base almacenada', 'balance': 'Saldo',
            }
            corte_labels = {
                'year_movement': 'Mov. año', 'final_balance': 'Saldo final',
            }

            columns.append(Command.create({
                'sequence': field.sequence,
                'field_name': field.name,
                'field_type': field.source,
                'field_type_display': source_label,
                'attribute': field.attribute or '',
                'xsd_type': field.xsd_type or '',
                'xsd_required': field.xsd_required,
                'xsd_status': xsd_status,
                'xsd_info': xsd_info,
                'max_length': field.max_length,
                'account_count': account_count,
                'nature': formula_labels.get(nature, nature),
                'format_field_id': field.id,
                'corte': corte_labels.get(field.corte, '') if field.source == 'journal_items' else '',
                'applies_smaller': field.applies_smaller_amounts,
                'principal_ref': field.principal_field_id.name or '',
                'default_value': field.default_value or '',
            }))

        # Validar campos requeridos XSD faltantes
        if xsd_attrs:
            configured_attrs = set(format_fields.mapped('attribute'))
            for attr_name, attr_info in xsd_attrs.items():
                if attr_name not in configured_attrs and attr_info.get('required'):
                    issues.append(Command.create({
                        'severity': 'error',
                        'category': 'xsd',
                        'message': f"Campo requerido XSD '{attr_name}' no configurado en el formato",
                    }))
                elif attr_name not in configured_attrs:
                    issues.append(Command.create({
                        'severity': 'warning',
                        'category': 'xsd',
                        'message': f"Campo opcional XSD '{attr_name}' no configurado",
                    }))

        # Validar configuración de líneas
        if not setting.format_setting_line_ids:
            issues.append(Command.create({
                'severity': 'error',
                'category': 'config',
                'message': 'No hay líneas de configuración. Use "Cargar columnas" primero.',
            }))

        # Validar conceptos
        total_concepts = 0
        total_accounts_all = 0
        if format_obj.apply_concepts:
            concepts = self.env['l10n_co.exogenous_concept'].search([
                ('format_id', '=', format_obj.id),
                ('active', '=', True)
            ])
            total_concepts = len(concepts)

            value_fields = format_fields.filtered(lambda f: f.source == 'journal_items')
            expected_lines = total_concepts * len(value_fields)
            actual_lines = len(setting.format_setting_line_ids)

            if actual_lines < expected_lines:
                issues.append(Command.create({
                    'severity': 'warning',
                    'category': 'config',
                    'message': f"Líneas configuradas: {actual_lines} de {expected_lines} esperadas ({total_concepts} conceptos x {len(value_fields)} campos)",
                }))

            for concept in concepts:
                if concept.field_account_ids:
                    for fa in concept.field_account_ids:
                        total_accounts_all += len(fa.account_ids)
                        if not fa.account_ids:
                            issues.append(Command.create({
                                'severity': 'warning',
                                'category': 'accounts',
                                'message': f"Concepto [{concept.code}] campo '{fa.format_field_id.name}': sin cuentas asignadas",
                            }))
                else:
                    issues.append(Command.create({
                        'severity': 'warning',
                        'category': 'accounts',
                        'message': f"Concepto [{concept.code}] {concept.name}: sin cuentas asignadas",
                    }))
        else:
            for field in format_fields.filtered(lambda f: f.source == 'journal_items'):
                total_accounts_all += field.account_count
                if field.account_count == 0:
                    issues.append(Command.create({
                        'severity': 'warning',
                        'category': 'accounts',
                        'message': f"Campo '{field.name}': sin cuentas asignadas",
                    }))

        # Validar columna principal
        value_fields = format_fields.filtered(lambda f: f.source == 'journal_items')
        for field in value_fields:
            if field.principal_field_id and field.principal_field_id not in value_fields:
                issues.append(Command.create({
                    'severity': 'error',
                    'category': 'config',
                    'message': f"Campo '{field.name}': referencia como principal a '{field.principal_field_id.name}' que no pertenece al formato",
                }))

        # Validar cuantía menor vs formato
        if format_obj.applying_smaller_amounts:
            smaller_fields = value_fields.filtered(lambda f: f.applies_smaller_amounts)
            if not smaller_fields:
                issues.append(Command.create({
                    'severity': 'warning',
                    'category': 'config',
                    'message': 'El formato aplica cuantías menores pero ninguna columna tiene "Aplica cuantía" activado',
                }))

        # Validar corte configurado en columnas de valor
        for field in value_fields:
            if not field.corte:
                issues.append(Command.create({
                    'severity': 'warning',
                    'category': 'config',
                    'message': f"Campo '{field.name}': sin corte configurado (Movimiento del año / Saldo final)",
                }))

        # Validar XML
        if not format_obj.xml_element_name:
            issues.append(Command.create({
                'severity': 'warning',
                'category': 'xml',
                'message': 'No se ha configurado el elemento XML del formato',
            }))

        # Validar fechas
        if format_obj.is_it_with_date_range:
            if not setting.date_start or not setting.date_end:
                issues.append(Command.create({
                    'severity': 'error',
                    'category': 'config',
                    'message': 'El formato requiere rango de fechas pero no están configuradas',
                }))

        # Calcular totales
        total_errors = len([i for i in issues if i[2].get('severity') == 'error'])
        total_warnings = len([i for i in issues if i[2].get('severity') == 'warning'])

        if total_errors > 0:
            audit_result = 'error'
        elif total_warnings > 0:
            audit_result = 'warning'
        else:
            audit_result = 'ok'

        self.write({
            'column_line_ids': columns,
            'issue_line_ids': issues,
            'total_columns': len(format_fields),
            'total_fixed': len(format_fields.filtered(lambda f: f.source == 'contact')),
            'total_value': len(format_fields.filtered(lambda f: f.source == 'journal_items')),
            'total_issues': total_errors,
            'total_warnings': total_warnings,
            'total_concepts': total_concepts,
            'total_accounts': total_accounts_all,
            'audit_result': audit_result,
        })

        return {
            'type': 'ir.actions.act_window',
            'name': _('Auditoría - %s') % format_obj.code,
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def _get_xsd_attributes(self):
        """Obtiene atributos del XSD"""
        xsd_doc = self.format_setting_id._get_xsd_document()
        if not xsd_doc:
            return {}

        attrs = {}
        try:
            root = xsd_doc.getroot()
            ns = {'xs': 'http://www.w3.org/2001/XMLSchema'}
            for attr in root.findall('.//xs:attribute', ns):
                attr_name = attr.get('name')
                if attr_name:
                    use = attr.get('use', 'optional')
                    attr_type = ''
                    simple_type = attr.find('xs:simpleType', ns)
                    if simple_type is not None:
                        restriction = simple_type.find('xs:restriction', ns)
                        if restriction is not None:
                            attr_type = restriction.get('base', '').replace('xs:', '')

                    attrs[attr_name] = {
                        'required': use == 'required',
                        'type': attr_type,
                    }
        except Exception as e:
            _logger.error("Error leyendo XSD: %s", str(e))

        return attrs

    def action_export_excel(self):
        """Exporta la auditoría a Excel"""
        self.ensure_one()

        wb = Workbook()

        # Hoja 1: Columnas del reporte
        ws1 = wb.active
        ws1.title = "Columnas Reporte"

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        error_fill = PatternFill(start_color='FF6B6B', end_color='FF6B6B', fill_type='solid')
        warning_fill = PatternFill(start_color='FFD93D', end_color='FFD93D', fill_type='solid')
        ok_fill = PatternFill(start_color='6BCB77', end_color='6BCB77', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        headers = ['Orden', 'Tipo', 'Nombre columna', 'Atributo XSD', 'Tipo XSD',
                   'Requerido', 'Long. Máx', 'No. Cuentas', 'Corte', 'Fórmula',
                   'Cuantía', 'Principal', 'Defecto', 'Estado XSD']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border

        col_widths = [8, 10, 35, 15, 12, 12, 12, 12, 12, 15, 10, 20, 10, 20]
        for i, width in enumerate(col_widths, 1):
            ws1.column_dimensions[get_column_letter(i)].width = width

        for idx, line in enumerate(self.column_line_ids, 2):
            ws1.cell(row=idx, column=1, value=line.sequence).border = thin_border
            ws1.cell(row=idx, column=2, value=line.field_type_display).border = thin_border
            ws1.cell(row=idx, column=3, value=line.field_name).border = thin_border
            ws1.cell(row=idx, column=4, value=line.attribute).border = thin_border
            ws1.cell(row=idx, column=5, value=line.xsd_type).border = thin_border
            ws1.cell(row=idx, column=6, value='Si' if line.xsd_required else 'No').border = thin_border
            ws1.cell(row=idx, column=7, value=line.max_length).border = thin_border
            ws1.cell(row=idx, column=8, value=line.account_count if line.field_type == 'journal_items' else '').border = thin_border
            ws1.cell(row=idx, column=9, value=line.corte or '').border = thin_border
            ws1.cell(row=idx, column=10, value=line.nature or '').border = thin_border
            ws1.cell(row=idx, column=11, value='Si' if line.applies_smaller else '').border = thin_border
            ws1.cell(row=idx, column=12, value=line.principal_ref or '').border = thin_border
            ws1.cell(row=idx, column=13, value=line.default_value or '').border = thin_border
            status_cell = ws1.cell(row=idx, column=14, value=line.xsd_info)
            status_cell.border = thin_border
            if line.xsd_status == 'ok':
                status_cell.fill = ok_fill
            elif line.xsd_status in ('extra', 'missing_attr'):
                status_cell.fill = warning_fill

        # Hoja 2: Problemas
        ws2 = wb.create_sheet("Problemas")
        headers2 = ['Severidad', 'Categoría', 'Descripción']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border

        ws2.column_dimensions['A'].width = 15
        ws2.column_dimensions['B'].width = 15
        ws2.column_dimensions['C'].width = 80

        category_labels = {
            'xsd': 'XSD', 'config': 'Configuración',
            'accounts': 'Cuentas', 'xml': 'XML'
        }

        for idx, issue in enumerate(self.issue_line_ids, 2):
            sev_cell = ws2.cell(row=idx, column=1, value='ERROR' if issue.severity == 'error' else 'Advertencia')
            sev_cell.border = thin_border
            sev_cell.fill = error_fill if issue.severity == 'error' else warning_fill
            ws2.cell(row=idx, column=2, value=category_labels.get(issue.category, issue.category)).border = thin_border
            ws2.cell(row=idx, column=3, value=issue.message).border = thin_border

        if not self.issue_line_ids:
            ws2.cell(row=2, column=1, value='Sin problemas')
            ws2.merge_cells('A2:C2')
            ws2.cell(row=2, column=1).fill = ok_fill

        output = io.BytesIO()
        wb.save(output)

        file_name = f"auditoria_{self.format_id.code}.xlsx"
        self.binary_file = base64.b64encode(output.getvalue())
        self.binary_file_name = file_name

        return {
            'type': 'ir.actions.act_window',
            'name': _('Auditoría - %s') % self.format_id.code,
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }


class L10nCoExogenousAuditWizardColumn(models.TransientModel):
    """Línea de columna en la auditoría"""
    _name = 'l10n_co.exogenous_audit_wizard.column'
    _description = 'Columna de auditoría'
    _order = 'sequence'

    wizard_id = fields.Many2one('l10n_co.exogenous_audit_wizard', required=True, ondelete='cascade')
    format_field_id = fields.Many2one('l10n_co.exogenous_format_field', string='Campo')
    sequence = fields.Integer(string='Orden')
    field_name = fields.Char(string='Nombre columna')
    field_type = fields.Char(string='Tipo interno')
    field_type_display = fields.Char(string='Tipo')
    attribute = fields.Char(string='Atributo XSD')
    xsd_type = fields.Char(string='Tipo XSD')
    xsd_required = fields.Boolean(string='Requerido')
    xsd_status = fields.Char(string='Estado XSD interno')
    xsd_info = fields.Char(string='Estado XSD')
    max_length = fields.Integer(string='Long. Máx')
    account_count = fields.Integer(string='No. Cuentas')
    nature = fields.Char(string='Fórmula')
    corte = fields.Char(string='Corte')
    applies_smaller = fields.Boolean(string='Aplica cuantía')
    principal_ref = fields.Char(string='Principal')
    default_value = fields.Char(string='Valor defecto')


class L10nCoExogenousAuditWizardIssue(models.TransientModel):
    """Problema detectado en la auditoría"""
    _name = 'l10n_co.exogenous_audit_wizard.issue'
    _description = 'Problema de auditoría'

    wizard_id = fields.Many2one('l10n_co.exogenous_audit_wizard', required=True, ondelete='cascade')
    severity = fields.Selection([
        ('error', 'Error'),
        ('warning', 'Advertencia'),
    ], string='Severidad')
    category = fields.Selection([
        ('xsd', 'XSD'),
        ('config', 'Configuración'),
        ('accounts', 'Cuentas'),
        ('xml', 'XML'),
    ], string='Categoría')
    message = fields.Char(string='Descripción')
