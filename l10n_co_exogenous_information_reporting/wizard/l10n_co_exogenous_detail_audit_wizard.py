# -*- coding: utf-8 -*-
import base64
import io
import logging
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from odoo import api, fields, models, _, Command
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class L10nCoExogenousDetailAuditWizard(models.TransientModel):
    """Wizard de auditoria auxiliar que muestra las lineas de movimiento
    tomadas para cada concepto/columna del reporte exogeno"""
    _name = 'l10n_co.exogenous_detail_audit_wizard'
    _description = 'Auditoría auxiliar detallada de exógena'

    format_setting_id = fields.Many2one(
        'l10n_co.exogenous_format_setting',
        string='Configuración',
        required=True)
    format_id = fields.Many2one(
        'l10n_co.exogenous_format',
        string='Formato',
        related='format_setting_id.format_id',
        readonly=True)
    company_id = fields.Many2one(
        'res.company',
        string='Compañía',
        related='format_setting_id.company_id',
        readonly=True)

    detail_line_ids = fields.One2many(
        'l10n_co.exogenous_detail_audit_wizard.line',
        'wizard_id',
        string='Detalle de líneas')
    summary_line_ids = fields.One2many(
        'l10n_co.exogenous_detail_audit_wizard.summary',
        'wizard_id',
        string='Resumen por concepto')

    total_lines = fields.Integer(string='Total líneas', readonly=True)
    total_debit = fields.Float(string='Total débito', readonly=True, digits=(16, 2))
    total_credit = fields.Float(string='Total crédito', readonly=True, digits=(16, 2))
    total_formula = fields.Float(string='Total fórmula', readonly=True, digits=(16, 2))

    binary_file = fields.Binary(string='Archivo Excel')
    binary_file_name = fields.Char(string='Nombre del archivo')

    def action_generate_detail(self):
        """Genera el detalle auxiliar con todas las lineas de movimiento"""
        self.ensure_one()
        self.detail_line_ids = [Command.clear()]
        self.summary_line_ids = [Command.clear()]

        setting = self.format_setting_id
        detail_lines = []
        summary_data = {}

        for setting_line in setting.format_setting_line_ids:
            field = setting_line.format_field_id
            concept = setting_line.concept_id if setting.apply_concepts else False

            concept_code = concept.code if concept else ''
            concept_name = concept.name if concept else field.name
            group_key = f"{concept_code}|{field.name}" if concept else field.name

            accounts = setting._get_accounts(setting_line)
            if not accounts:
                continue

            partner_ids, move_lines = setting._get_information_by_account_move_line(
                accounts.ids, format_field=field, setting_line=setting_line)

            if not move_lines:
                if group_key not in summary_data:
                    summary_data[group_key] = {
                        'concept_code': concept_code,
                        'concept_name': concept_name,
                        'field_name': field.name,
                        'formula': field.nature_account or 'db_cr',
                        'total_debit': 0, 'total_credit': 0,
                        'total_formula': 0, 'line_count': 0,
                    }
                continue

            formula = field.nature_account or 'db_cr'

            for aml in move_lines:
                debit = aml.get('debit', 0) or 0
                credit = aml.get('credit', 0) or 0
                formula_value = setting._compute_formula(aml, formula)

                account_info = aml.get('account_id', [False, ''])
                partner_info = aml.get('partner_id', [False, ''])
                move_info = aml.get('move_id', [False, ''])
                aml_date = aml.get('date', False)

                detail_lines.append(Command.create({
                    'concept_code': concept_code,
                    'concept_name': concept_name[:100] if concept_name else '',
                    'field_name': field.name,
                    'account_id': account_info[0] if isinstance(account_info, (list, tuple)) else account_info,
                    'account_display': account_info[1] if isinstance(account_info, (list, tuple)) else str(account_info),
                    'partner_id': partner_info[0] if isinstance(partner_info, (list, tuple)) else partner_info,
                    'partner_display': partner_info[1] if isinstance(partner_info, (list, tuple)) else str(partner_info),
                    'move_display': move_info[1] if isinstance(move_info, (list, tuple)) else str(move_info),
                    'date': aml_date,
                    'debit': debit,
                    'credit': credit,
                    'formula_value': formula_value,
                    'formula_type': formula,
                }))

                if group_key not in summary_data:
                    summary_data[group_key] = {
                        'concept_code': concept_code,
                        'concept_name': concept_name,
                        'field_name': field.name,
                        'formula': formula,
                        'total_debit': 0, 'total_credit': 0,
                        'total_formula': 0, 'line_count': 0,
                    }

                summary_data[group_key]['total_debit'] += debit
                summary_data[group_key]['total_credit'] += credit
                summary_data[group_key]['total_formula'] += formula_value
                summary_data[group_key]['line_count'] += 1

        summary_lines = []
        for data in summary_data.values():
            summary_lines.append((0, 0, data))

        grand_total_debit = sum(d['total_debit'] for d in summary_data.values())
        grand_total_credit = sum(d['total_credit'] for d in summary_data.values())
        grand_total_formula = sum(d['total_formula'] for d in summary_data.values())

        self.write({
            'detail_line_ids': detail_lines,
            'summary_line_ids': summary_lines,
            'total_lines': len(detail_lines),
            'total_debit': grand_total_debit,
            'total_credit': grand_total_credit,
            'total_formula': grand_total_formula,
        })

        return {
            'type': 'ir.actions.act_window',
            'name': _('Auxiliar detallado - %s') % self.format_id.code,
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
            'context': {'dialog_size': 'extra-large'},
        }

    def action_export_excel(self):
        """Exporta el auxiliar detallado a Excel"""
        self.ensure_one()

        wb = Workbook()
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        summary_fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
        total_font = Font(bold=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'))

        # Hoja 1: Resumen por concepto
        ws1 = wb.active
        ws1.title = "Resumen"

        headers1 = ['Concepto', 'Nombre', 'Columna', 'Fórmula',
                     'Total débito', 'Total crédito', 'Total fórmula', 'Líneas']
        for col, header in enumerate(headers1, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border

        widths1 = [12, 40, 30, 15, 18, 18, 18, 10]
        for i, w in enumerate(widths1, 1):
            ws1.column_dimensions[get_column_letter(i)].width = w

        formula_labels = {
            'debit': 'DB', 'credit': 'CR', 'db_cr': 'DB-CR', 'cr_db': 'CR-DB',
            'base_calc_db': 'Base DB', 'base_calc_cr': 'Base CR',
            'base_calc_db_cr': 'Base DB-CR', 'base_calc_cr_db': 'Base CR-DB',
            'tax_base_amount': 'Base almacenada', 'balance': 'Saldo',
        }

        for idx, line in enumerate(self.summary_line_ids, 2):
            ws1.cell(row=idx, column=1, value=line.concept_code).border = thin_border
            ws1.cell(row=idx, column=2, value=line.concept_name[:60] if line.concept_name else '').border = thin_border
            ws1.cell(row=idx, column=3, value=line.field_name).border = thin_border
            ws1.cell(row=idx, column=4, value=formula_labels.get(line.formula, line.formula)).border = thin_border
            ws1.cell(row=idx, column=5, value=round(line.total_debit, 2)).border = thin_border
            ws1.cell(row=idx, column=6, value=round(line.total_credit, 2)).border = thin_border
            ws1.cell(row=idx, column=7, value=round(line.total_formula, 2)).border = thin_border
            ws1.cell(row=idx, column=8, value=line.line_count).border = thin_border

        total_row = len(self.summary_line_ids) + 2
        ws1.cell(row=total_row, column=1, value='TOTAL').font = total_font
        ws1.cell(row=total_row, column=1).fill = summary_fill
        ws1.cell(row=total_row, column=5, value=round(self.total_debit, 2)).font = total_font
        ws1.cell(row=total_row, column=5).fill = summary_fill
        ws1.cell(row=total_row, column=6, value=round(self.total_credit, 2)).font = total_font
        ws1.cell(row=total_row, column=6).fill = summary_fill
        ws1.cell(row=total_row, column=7, value=round(self.total_formula, 2)).font = total_font
        ws1.cell(row=total_row, column=7).fill = summary_fill
        ws1.cell(row=total_row, column=8, value=self.total_lines).font = total_font
        ws1.cell(row=total_row, column=8).fill = summary_fill

        # Hoja 2: Detalle auxiliar
        ws2 = wb.create_sheet("Auxiliar detallado")

        headers2 = ['Concepto', 'Columna', 'Cuenta', 'Tercero', 'Fecha',
                     'Asiento', 'Débito', 'Crédito', 'Fórmula', 'Valor']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border

        widths2 = [12, 25, 30, 35, 12, 20, 16, 16, 12, 16]
        for i, w in enumerate(widths2, 1):
            ws2.column_dimensions[get_column_letter(i)].width = w

        for idx, line in enumerate(self.detail_line_ids, 2):
            ws2.cell(row=idx, column=1, value=line.concept_code).border = thin_border
            ws2.cell(row=idx, column=2, value=line.field_name).border = thin_border
            ws2.cell(row=idx, column=3, value=line.account_display).border = thin_border
            ws2.cell(row=idx, column=4, value=line.partner_display).border = thin_border
            ws2.cell(row=idx, column=5, value=str(line.date) if line.date else '').border = thin_border
            ws2.cell(row=idx, column=6, value=line.move_display).border = thin_border
            ws2.cell(row=idx, column=7, value=round(line.debit, 2)).border = thin_border
            ws2.cell(row=idx, column=8, value=round(line.credit, 2)).border = thin_border
            ws2.cell(row=idx, column=9, value=formula_labels.get(line.formula_type, line.formula_type)).border = thin_border
            ws2.cell(row=idx, column=10, value=round(line.formula_value, 2)).border = thin_border

        output = io.BytesIO()
        wb.save(output)

        file_name = f"auxiliar_{self.format_id.code}.xlsx"
        self.binary_file = base64.b64encode(output.getvalue())
        self.binary_file_name = file_name

        return {
            'type': 'ir.actions.act_window',
            'name': _('Auxiliar detallado - %s') % self.format_id.code,
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }


class L10nCoExogenousDetailAuditLine(models.TransientModel):
    """Linea de detalle del auxiliar"""
    _name = 'l10n_co.exogenous_detail_audit_wizard.line'
    _description = 'Línea de detalle auxiliar'
    _order = 'concept_code, field_name, date'

    wizard_id = fields.Many2one(
        'l10n_co.exogenous_detail_audit_wizard',
        required=True, ondelete='cascade')
    concept_code = fields.Char(string='Concepto')
    concept_name = fields.Char(string='Nombre concepto')
    field_name = fields.Char(string='Columna')
    account_id = fields.Integer(string='ID cuenta')
    account_display = fields.Char(string='Cuenta')
    partner_id = fields.Integer(string='ID tercero')
    partner_display = fields.Char(string='Tercero')
    move_display = fields.Char(string='Asiento')
    date = fields.Date(string='Fecha')
    debit = fields.Float(string='Débito', digits=(16, 2))
    credit = fields.Float(string='Crédito', digits=(16, 2))
    formula_value = fields.Float(string='Valor fórmula', digits=(16, 2))
    formula_type = fields.Char(string='Fórmula')


class L10nCoExogenousDetailAuditSummary(models.TransientModel):
    """Resumen por concepto del auxiliar"""
    _name = 'l10n_co.exogenous_detail_audit_wizard.summary'
    _description = 'Resumen por concepto auxiliar'
    _order = 'concept_code, field_name'

    wizard_id = fields.Many2one(
        'l10n_co.exogenous_detail_audit_wizard',
        required=True, ondelete='cascade')
    concept_code = fields.Char(string='Concepto')
    concept_name = fields.Char(string='Nombre')
    field_name = fields.Char(string='Columna')
    formula = fields.Char(string='Fórmula')
    total_debit = fields.Float(string='Total débito', digits=(16, 2))
    total_credit = fields.Float(string='Total crédito', digits=(16, 2))
    total_formula = fields.Float(string='Total fórmula', digits=(16, 2))
    line_count = fields.Integer(string='Líneas')
